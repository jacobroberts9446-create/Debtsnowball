"""Create README screenshots from real workflows and synthetic data."""

from __future__ import annotations

import tempfile
from collections.abc import Callable
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont

from app.config import Config
from app.excel_writer import ExcelWriter
from app.forecast_engine import ForecastEngine
from app.history import PlanHistoryService
from app.preferences import RecentPlanPreferences
from app.workflows.plan_progress import run_plan_progress_menu, show_progress_summary
from app.workflows.saved_plans import run_saved_plans_menu
from run import run_main_menu

PROJECT_ROOT = Path(__file__).resolve().parent.parent
IMAGE_DIR = PROJECT_ROOT / "docs" / "images"
FONT_PATH = Path("C:/Windows/Fonts/consola.ttf")
BOLD_FONT_PATH = Path("C:/Windows/Fonts/consolab.ttf")


def main() -> int:
    """Generate current branded screenshots using only synthetic values."""
    IMAGE_DIR.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="debtpilot-screenshots-") as temporary:
        temp_dir = Path(temporary)
        config = synthetic_config()
        forecast = ForecastEngine(config).forecast()
        service = PlanHistoryService(temp_dir / "plans.sqlite")
        saved = service.save_generated_plan(
            name="Sample Household Plan",
            config=config,
            forecast=forecast,
            starting_savings=config.settings.starting_savings,
            starting_debts=config.debts,
            description="Synthetic example data",
        )
        preferences = RecentPlanPreferences(temp_dir / "preferences.json")

        render_console("Home", capture_home(), IMAGE_DIR / "home.png")
        render_console(
            "Saved Plans",
            capture_saved_plans(service, preferences),
            IMAGE_DIR / "saved_plans.png",
        )
        render_console(
            "Track Progress",
            capture_progress_menu(service, saved.plan),
            IMAGE_DIR / "track_progress.png",
        )
        render_console(
            "Progress Summary",
            capture_progress_summary(service, saved.plan),
            IMAGE_DIR / "progress_summary.png",
        )

        workbook_path = temp_dir / "debtpilot_plan.xlsx"
        ExcelWriter(workbook_path).write([], forecast)
        render_workbook(workbook_path, IMAGE_DIR / "workbook.png")
    return 0


def synthetic_config() -> Config:
    """Return a compact configuration containing no user financial data."""
    return Config().load_mapping(
        {
            "budget": {
                "paycheck": "1800.00",
                "first_paycheck": "2026-08-07",
                "rent_per_paycheck": "450.00",
                "insurance_per_paycheck": "75.00",
                "personal_per_paycheck": "250.00",
                "starting_savings": "1200.00",
                "savings_goal": "3000.00",
                "snowball_split": "0.50",
            },
            "bills": [{"name": "Utilities", "amount": "180.00", "due_day": 15}],
            "debts": [
                {
                    "name": "Credit Card",
                    "balance": "3200.00",
                    "apr": "18.90",
                    "minimum": "95.00",
                    "due_day": 20,
                    "snowball_order": 1,
                },
                {
                    "name": "Auto Loan",
                    "balance": "8400.00",
                    "apr": "5.25",
                    "minimum": "310.00",
                    "due_day": 25,
                    "snowball_order": 2,
                },
            ],
        }
    )


def capture_home() -> list[str]:
    """Capture the current first-launch menu."""
    output = []
    run_main_menu(
        input_func=canned_input(["4"]),
        output_func=output.append,
    )
    return [line for line in output if line != "Success: Goodbye."]


def capture_saved_plans(
    service: PlanHistoryService,
    preferences: RecentPlanPreferences,
) -> list[str]:
    """Capture the saved-plan selection screen."""
    output = []
    run_saved_plans_menu(
        plan_history_service_factory=lambda: service,
        preferences_factory=lambda: preferences,
        input_func=canned_input(["2"]),
        output_func=output.append,
    )
    return output


def capture_progress_menu(service: PlanHistoryService, plan) -> list[str]:
    """Capture the selected plan's progress menu."""
    output = []
    run_plan_progress_menu(
        service,
        plan,
        input_func=canned_input(["6"]),
        output_func=output.append,
    )
    return output


def capture_progress_summary(service: PlanHistoryService, plan) -> list[str]:
    """Capture the read-only progress summary."""
    output = []
    show_progress_summary(
        service,
        plan,
        input_func=canned_input([""]),
        output_func=output.append,
    )
    return output


def canned_input(values: list[str]) -> Callable[[str], str]:
    """Return an input function that consumes deterministic responses."""
    responses = iter(values)
    return lambda _prompt: next(responses)


def render_console(title: str, lines: list[str], destination: Path) -> None:
    """Render captured plain-text output in a terminal-style frame."""
    font = ImageFont.truetype(str(FONT_PATH), 21)
    title_font = ImageFont.truetype(str(BOLD_FONT_PATH), 16)
    visible_lines = _trim_blank_edges(lines)
    width = 1280
    line_height = 29
    height = max(460, 68 + line_height * len(visible_lines) + 28)
    image = Image.new("RGB", (width, height), "#0c1117")
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, width, 44), fill="#202a35")
    draw.ellipse((16, 15, 28, 27), fill="#ff5f57")
    draw.ellipse((36, 15, 48, 27), fill="#febc2e")
    draw.ellipse((56, 15, 68, 27), fill="#28c840")
    draw.text((88, 12), f"DebtPilot - {title}", font=title_font, fill="#dbe7f2")
    y = 61
    for line in visible_lines:
        color = "#58d6c7" if "DebtPilot v" in line else "#e6edf3"
        if line.startswith("Warning:"):
            color = "#f2cc60"
        draw.text((26, y), line, font=font, fill=color)
        y += line_height
    image.save(destination)


def render_workbook(workbook_path: Path, destination: Path) -> None:
    """Render actual Dashboard worksheet values as a workbook preview."""
    workbook = load_workbook(workbook_path, data_only=True)
    try:
        sheet = workbook["Dashboard"]
        rows = [
            [_display_cell_value(sheet.cell(row=row, column=column)) for column in range(1, 3)]
            for row in range(1, min(sheet.max_row, 16) + 1)
        ]
    finally:
        workbook.close()

    width, height = 1280, 720
    image = Image.new("RGB", (width, height), "#f3f6f8")
    draw = ImageDraw.Draw(image)
    title_font = ImageFont.truetype(str(BOLD_FONT_PATH), 24)
    cell_font = ImageFont.truetype(str(FONT_PATH), 18)
    draw.rectangle((0, 0, width, 52), fill="#217346")
    draw.text((22, 13), "DebtPilot - Dashboard", font=title_font, fill="white")
    left, top = 42, 84
    column_widths = [430, 360]
    row_height = 36
    for row_index, values in enumerate(rows):
        y = top + row_index * row_height
        for column_index, value in enumerate(values):
            x = left + sum(column_widths[:column_index])
            fill = "#d9e9f5" if row_index in {0, 2} else "white"
            draw.rectangle(
                (x, y, x + column_widths[column_index], y + row_height),
                fill=fill,
                outline="#b8c2cc",
            )
            text = "" if value is None else str(value)
            draw.text((x + 10, y + 7), text, font=cell_font, fill="#17202a")
    draw.text(
        (840, 104),
        "Workbook preview",
        font=title_font,
        fill="#217346",
    )
    draw.text(
        (840, 150),
        "Dashboard\nPay Periods\nDebt Tracking\nSavings Progress\nForecast Charts",
        font=cell_font,
        fill="#34495e",
        spacing=13,
    )
    image.save(destination)


def _trim_blank_edges(lines: list[str]) -> list[str]:
    """Remove only leading and trailing blank captured output."""
    result = list(lines)
    while result and not result[0]:
        result.pop(0)
    while result and not result[-1]:
        result.pop()
    return result


def _display_cell_value(cell) -> str:
    """Format one workbook cell the way a user sees it in Excel."""
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return f"{value:%b} {value.day}, {value:%Y}"
    if "$" in cell.number_format and isinstance(value, (int, float, Decimal)):
        return f"${value:,.2f}"
    return str(value)


if __name__ == "__main__":
    raise SystemExit(main())
