"""
excel_writer.py

Writes budget summaries to an Excel workbook.
"""

from dataclasses import asdict
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.budget_engine import PayPeriodSummary
from app.money import ZERO_MONEY, excel_number, money
from app.models import (
    AssumptionDifference,
    ForecastActualComparison,
    ForecastActualPeriodComparison,
    DebtFreeTargetResult,
    DebtFreeTargetStatus,
    ForecastSummary,
    Plan,
    PlanComparison,
    PlanVersion,
    ScenarioComparison,
    ScenarioResult,
)
from app.paths import default_workbook_path


class ExcelWriter:
    """Creates a simple workbook from PayPeriodSummary objects."""

    def __init__(self, filename: str | Path | None = None) -> None:
        self.path = default_workbook_path() if filename is None else Path(filename)

    def write(
        self,
        summaries: list[PayPeriodSummary],
        forecast: ForecastSummary | None = None,
        scenario_comparison: ScenarioComparison | None = None,
        target_result: DebtFreeTargetResult | None = None,
    ) -> Path:
        """Write pay-period, debt, payoff, and savings worksheets."""
        self.path.parent.mkdir(parents=True, exist_ok=True)

        workbook = Workbook()
        dashboard_sheet = workbook.active
        dashboard_sheet.title = "Dashboard"

        pay_period_sheet = workbook.create_sheet("Pay Period Summaries")

        self._write_pay_period_summaries(pay_period_sheet, summaries)
        self._write_active_debts(workbook.create_sheet("Active Debts"), summaries)
        self._write_paid_off_debts(workbook.create_sheet("Paid-Off Debts"), summaries)
        self._write_savings_progress(
            workbook.create_sheet("Savings Progress"),
            summaries,
            forecast,
        )
        self._write_forecast(workbook.create_sheet("Forecast"), forecast)
        self._write_scenario_comparison(
            workbook.create_sheet("Scenario Comparison"),
            scenario_comparison,
        )
        if target_result is not None:
            self._write_debt_free_target(
                workbook.create_sheet("Debt-Free Target"),
                target_result,
            )
        self._write_dashboard(
            dashboard_sheet,
            summaries,
            forecast,
            scenario_comparison,
            target_result,
        )

        workbook.save(self.path)
        return self.path

    def write_history_report(
        self,
        plan: Plan,
        versions: list[PlanVersion],
        comparison: PlanComparison | None = None,
        actual_comparison: ForecastActualComparison | None = None,
        forecast_snapshots: list[dict[str, object]] | None = None,
        actual_periods: list[ForecastActualPeriodComparison] | None = None,
        debt_history: list[dict[str, object]] | None = None,
        savings_history: list[dict[str, object]] | None = None,
        warnings: list[dict[str, object]] | None = None,
        assumption_differences: list[AssumptionDifference] | None = None,
    ) -> Path:
        """Write optional saved-history sheets when history data is requested."""
        self.path.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        history_sheet = workbook.active
        history_sheet.title = "Plan History"
        self._write_plan_history(history_sheet, plan, versions, forecast_snapshots)
        if comparison is not None:
            self._write_plan_comparison(workbook.create_sheet("Plan Comparison"), comparison)
        if actual_comparison is not None:
            self._write_forecast_actual(
                workbook.create_sheet("Forecast vs Actual"),
                actual_comparison,
            )
        if actual_periods is not None:
            self._write_history_dict_sheet(
                workbook.create_sheet("Period Details"),
                [asdict(row) for row in actual_periods],
            )
        if debt_history is not None:
            self._write_history_dict_sheet(
                workbook.create_sheet("Debt History"),
                debt_history,
            )
        if savings_history is not None:
            self._write_history_dict_sheet(
                workbook.create_sheet("Savings History"),
                savings_history,
            )
        if warnings is not None:
            self._write_history_dict_sheet(
                workbook.create_sheet("History Warnings"),
                warnings,
            )
        if assumption_differences is not None:
            self._write_history_dict_sheet(
                workbook.create_sheet("Assumption Changes"),
                [asdict(row) for row in assumption_differences],
            )
        workbook.save(self.path)
        return self.path

    def _write_plan_history(
        self,
        sheet: Worksheet,
        plan: Plan,
        versions: list[PlanVersion],
        forecast_snapshots: list[dict[str, object]] | None = None,
    ) -> None:
        sheet["A1"] = f"Plan History: {plan.name}"
        sheet["A1"].font = Font(bold=True, size=14)
        snapshots_by_version = {
            int(snapshot["plan_version_id"]): snapshot
            for snapshot in forecast_snapshots or []
        }
        self._append_row(
            sheet,
            [
                "Version",
                "Created Date",
                "Change Note",
                "Configuration Fingerprint",
                "Forecast Fingerprint",
                "Debt-Free Date",
                "Total Projected Interest",
                "Starting Debt",
                "Ending Savings",
                "Warning Count",
                "Active",
                "Application Version",
                "Engine Version",
            ],
        )
        for version in versions:
            snapshot = snapshots_by_version.get(version.id, {})
            self._append_row(
                sheet,
                [
                    version.version_number,
                    version.created_at,
                    version.change_note,
                    version.config_fingerprint,
                    snapshot.get("forecast_fingerprint"),
                    self._date_cell(snapshot.get("debt_free_date")),
                    self._money_cell(snapshot.get("total_projected_interest")),
                    self._money_cell(snapshot.get("starting_debt")),
                    self._money_cell(snapshot.get("ending_savings")),
                    snapshot.get("warning_count"),
                    self._yes_no(version.active),
                    version.application_version,
                    version.forecast_engine_version,
                ],
            )
        self._format_table(sheet, currency_columns=[7, 8, 9], date_columns=[2, 6], header_row=2)

    def _write_plan_comparison(
        self,
        sheet: Worksheet,
        comparison: PlanComparison,
    ) -> None:
        sheet["A1"] = "Plan Comparison"
        sheet["A1"].font = Font(bold=True, size=14)
        rows = [
            ("Debt-Free Date Difference Days", comparison.debt_free_date_difference_days),
            ("Interest Difference", comparison.interest_difference),
            ("Debt Payment Difference", comparison.debt_payment_difference),
            ("Savings Difference", comparison.savings_difference),
            ("Personal Spending Difference", comparison.personal_spending_difference),
            ("Pay Period Difference", comparison.pay_period_difference),
            ("First Different Period", comparison.first_different_period),
            ("Payoff Order Changed", self._yes_no(comparison.payoff_order_changed)),
            ("Deadline Priority Changed", self._yes_no(comparison.deadline_priority_changed)),
            ("Feasible", self._yes_no(comparison.feasible)),
            ("Interpretation", comparison.explanation),
        ]
        self._append_row(sheet, ["Metric", "Value"])
        for row in rows:
            self._append_row(sheet, list(row))
        self._format_table(sheet, currency_columns=[2], date_columns=[2], header_row=2)

    def _write_forecast_actual(
        self,
        sheet: Worksheet,
        comparison: ForecastActualComparison,
    ) -> None:
        sheet["A1"] = "Forecast vs Actual"
        sheet["A1"].font = Font(bold=True, size=14)
        self._append_row(sheet, ["Metric", "Planned", "Actual", "Variance", "Status"])
        rows = [
            ("Income", comparison.planned_income, comparison.actual_income),
            ("Bills", comparison.planned_bills, comparison.actual_bills),
            (
                "Debt Payments",
                comparison.planned_debt_payments,
                comparison.actual_debt_payments,
            ),
            ("Savings", comparison.planned_savings, comparison.actual_savings),
            (
                "Personal Spending",
                comparison.planned_personal_spending,
                comparison.actual_personal_spending,
            ),
            (
                "Remaining Cash",
                comparison.planned_remaining_cash,
                comparison.actual_remaining_cash,
            ),
        ]
        for metric, planned, actual in rows:
            self._append_row(
                sheet,
                [metric, planned, actual, money(actual - planned), comparison.status],
            )
        self._format_table(sheet, currency_columns=[2, 3, 4], header_row=2)

    def _write_history_dict_sheet(
        self,
        sheet: Worksheet,
        rows: list[dict[str, object]],
    ) -> None:
        """Write a simple formatted sheet from persisted history rows."""
        if not rows:
            self._append_row(sheet, ["Status"])
            self._append_row(sheet, ["No rows available"])
            self._format_table(sheet, header_row=1)
            return

        headers = list(rows[0])
        self._append_row(sheet, headers)
        for row in rows:
            self._append_row(sheet, [row.get(header) for header in headers])
        currency_columns = [
            index
            for index, header in enumerate(headers, start=1)
            if any(
                marker in header
                for marker in (
                    "amount",
                    "balance",
                    "payment",
                    "interest",
                    "income",
                    "expense",
                    "savings",
                    "snowball",
                    "principal",
                    "withdrawal",
                    "cash",
                    "variance",
                    "contribution",
                    "debt",
                )
            )
        ]
        date_columns = [
            index
            for index, header in enumerate(headers, start=1)
            if "date" in header or header in {"pay_date", "created_at", "matched_at"}
        ]
        self._format_table(
            sheet,
            currency_columns=currency_columns,
            date_columns=date_columns,
        )

    def _write_pay_period_summaries(
        self,
        sheet: Worksheet,
        summaries: list[PayPeriodSummary],
    ) -> None:
        self._append_row(
            sheet,
            [
                "Pay Date",
                "Start Date",
                "End Date",
                "Income",
                "Bills Paid",
                "Required Fixed Expenses",
                "Normal Personal Allowance",
                "Actual Personal Allowance",
                "Debt Minimums",
                "Savings Goal",
                "Available After Required Payments",
                "Normal Savings",
                "Required Savings",
                "Savings Deposit",
                "Snowball Before Adjustment",
                "Snowball Redirected To Savings",
                "Personal Expense Reduction",
                "Snowball Payment",
                "Savings Shortfall",
                "Remaining Cash",
                "Active Debt Total",
                "Paid-Off Debt Count",
                "Chart Label",
            ]
        )

        for summary in summaries:
            active_debt_total = self._active_debt_total(summary)
            self._append_row(
                sheet,
                [
                    summary.pay_date,
                    summary.start_date,
                    summary.end_date,
                    summary.income,
                    summary.bills_paid,
                    summary.required_fixed_expenses,
                    summary.normal_personal_allowance,
                    summary.actual_personal_allowance,
                    summary.debt_minimums,
                    summary.active_savings_goal_name,
                    summary.available_after_required_payments,
                    summary.normal_savings_contribution,
                    summary.deadline_required_savings_contribution,
                    summary.savings_contribution,
                    summary.snowball_before_savings_adjustment,
                    summary.snowball_reduction,
                    summary.personal_expense_reduction,
                    summary.snowball_payment,
                    summary.projected_savings_shortfall,
                    summary.remaining_cash,
                    active_debt_total,
                    len(summary.paid_off_debts),
                    self._chart_date_label(summary.pay_date, len(summaries)),
                ]
            )

        self._format_table(
            sheet,
            currency_columns=[
                4,
                5,
                6,
                7,
                8,
                9,
                11,
                12,
                13,
                14,
                15,
                16,
                17,
                18,
                19,
                20,
                21,
            ],
            date_columns=[1, 2, 3],
        )

    def _write_active_debts(
        self,
        sheet: Worksheet,
        summaries: list[PayPeriodSummary],
    ) -> None:
        self._append_row(sheet, ["Pay Date", "Debt", "Balance", "Minimum", "Status"])

        for summary in summaries:
            for debt in summary.active_debt_balances:
                self._append_row(
                    sheet,
                    [
                        summary.pay_date,
                        debt.name,
                        debt.balance,
                        debt.minimum,
                        debt.status,
                    ]
                )

        self._format_table(sheet, currency_columns=[3, 4], date_columns=[1])

    def _write_paid_off_debts(
        self,
        sheet: Worksheet,
        summaries: list[PayPeriodSummary],
    ) -> None:
        self._append_row(sheet, ["Pay Date", "Debt", "Balance", "Minimum", "Status"])

        seen = set()
        for summary in summaries:
            for debt in summary.paid_off_debts:
                if debt.name in seen:
                    continue

                seen.add(debt.name)
                self._append_row(
                    sheet,
                    [
                        summary.pay_date,
                        debt.name,
                        debt.balance,
                        debt.minimum,
                        debt.status,
                    ]
                )

        self._format_table(sheet, currency_columns=[3, 4], date_columns=[1])

    def _write_savings_progress(
        self,
        sheet: Worksheet,
        summaries: list[PayPeriodSummary],
        forecast: ForecastSummary | None = None,
    ) -> None:
        if forecast is not None and (
            forecast.savings_stage_results or forecast.planned_withdrawal_results
        ):
            self._write_savings_plan_sections(sheet, forecast, summaries)
            return

        self._append_row(
            sheet,
            [
                "Pay Date",
                "Savings Deposit",
                "Savings Balance",
                "Savings Goal",
                "Remaining To Goal",
                "Chart Label",
            ]
        )

        for summary in summaries:
            remaining_to_goal = max(
                summary.savings_goal - summary.savings_balance,
                ZERO_MONEY,
            )
            self._append_row(
                sheet,
                [
                    summary.pay_date,
                    summary.savings_contribution,
                    summary.savings_balance,
                    summary.savings_goal,
                    remaining_to_goal,
                    self._chart_date_label(summary.pay_date, len(summaries)),
                ]
            )

        self._format_table(sheet, currency_columns=[2, 3, 4, 5], date_columns=[1])

    def _write_savings_plan_sections(
        self,
        sheet: Worksheet,
        forecast: ForecastSummary,
        summaries: list[PayPeriodSummary],
    ) -> None:
        sheet["A1"] = "Savings Plan Summary"
        sheet["A1"].font = Font(bold=True, size=16)
        summary_headers = [
            "Goal Name",
            "Start Date",
            "Target Date",
            "Target Amount",
            "Starting Balance",
            "Amount Needed",
            "Eligible Paychecks Remaining",
            "Projected Available Contributions",
            "Projected Balance At Deadline",
            "Projected Shortfall",
            "Additional Funding Needed",
            "Feasible",
            "Achieved Date",
            "Amount at Deadline",
            "Shortfall at Deadline",
            "Current/Ending Balance",
            "Status",
            "Days Early/Late",
        ]
        self._write_header(sheet, 3, summary_headers)
        for row_index, stage in enumerate(forecast.savings_stage_results, start=4):
            self._append_row(
                sheet,
                [
                    stage.goal_name,
                    stage.start_date,
                    stage.target_date,
                    self._cell_value(stage.target_amount),
                    self._cell_value(stage.starting_balance),
                    self._cell_value(stage.amount_needed),
                    stage.eligible_paychecks_remaining,
                    self._cell_value(stage.projected_available_contributions),
                    self._cell_value(stage.projected_balance_at_deadline),
                    self._cell_value(stage.projected_shortfall),
                    self._cell_value(stage.additional_funding_needed),
                    self._yes_no(stage.feasible_under_current_plan),
                    stage.achieved_date,
                    self._cell_value(stage.amount_at_deadline),
                    self._cell_value(stage.shortfall_at_deadline),
                    self._cell_value(stage.ending_balance),
                    stage.status.value,
                    stage.days_early_or_late,
                ]
            )

        withdrawals_title_row = 5 + len(forecast.savings_stage_results)
        sheet.cell(row=withdrawals_title_row, column=1, value="Planned Withdrawals")
        sheet.cell(row=withdrawals_title_row, column=1).font = Font(bold=True)
        withdrawal_header_row = withdrawals_title_row + 1
        withdrawal_headers = [
            "Name",
            "Scheduled Date",
            "Requested Amount",
            "Drain Balance",
            "Actual Amount Withdrawn",
            "Balance Before",
            "Balance After",
            "Applied Date",
            "Status",
        ]
        self._write_header(sheet, withdrawal_header_row, withdrawal_headers)
        for row_index, withdrawal in enumerate(
            forecast.planned_withdrawal_results,
            start=withdrawal_header_row + 1,
        ):
            self._append_row(
                sheet,
                [
                    withdrawal.name,
                    withdrawal.scheduled_date,
                    self._cell_value(withdrawal.requested_amount),
                    withdrawal.drain_balance,
                    self._cell_value(withdrawal.actual_amount_withdrawn),
                    self._cell_value(withdrawal.balance_before),
                    self._cell_value(withdrawal.balance_after),
                    withdrawal.applied_date,
                    withdrawal.status,
                ]
            )

        detail_title_row = withdrawal_header_row + max(
            len(forecast.planned_withdrawal_results),
            1,
        ) + 3
        sheet.cell(row=detail_title_row, column=1, value="Savings Progress Detail")
        sheet.cell(row=detail_title_row, column=1).font = Font(bold=True)
        detail_header_row = detail_title_row + 1
        detail_headers = [
            "Pay Date",
            "Active Goal",
            "Required Fixed Expenses",
            "Normal Personal Allowance",
            "Actual Personal Allowance",
            "Available After Required Payments",
            "Normal Savings",
            "Required Savings",
            "Savings Contribution",
            "Snowball Before Adjustment",
            "Snowball Redirected To Savings",
            "Personal Expense Reduction",
            "Actual Snowball Payment",
            "Projected Savings Shortfall",
            "Withdrawal",
            "Ending Savings Balance",
            "Active Target",
            "Progress",
            "Chart Label",
        ]
        self._write_header(sheet, detail_header_row, detail_headers)
        for row_index, summary in enumerate(summaries, start=detail_header_row + 1):
            self._append_row(
                sheet,
                [
                    summary.pay_date,
                    summary.active_savings_goal_name,
                    summary.required_fixed_expenses,
                    summary.normal_personal_allowance,
                    summary.actual_personal_allowance,
                    summary.available_after_required_payments,
                    summary.normal_savings_contribution,
                    summary.deadline_required_savings_contribution,
                    summary.savings_contribution,
                    summary.snowball_before_savings_adjustment,
                    summary.snowball_reduction,
                    summary.personal_expense_reduction,
                    summary.snowball_payment,
                    summary.projected_savings_shortfall,
                    summary.planned_withdrawal_amount,
                    summary.savings_balance,
                    summary.active_savings_target,
                    summary.goal_progress_percentage,
                    self._chart_date_label(summary.pay_date, len(summaries)),
                ]
            )

        self._format_savings_plan_sheet(
            sheet,
            detail_header_row=detail_header_row,
        )

    def _format_savings_plan_sheet(
        self,
        sheet: Worksheet,
        detail_header_row: int,
    ) -> None:
        non_currency_headers = {
            "Goal Name",
            "Start Date",
            "Target Date",
            "Eligible Paychecks Remaining",
            "Feasible",
            "Achieved Date",
            "Status",
            "Days Early/Late",
            "Name",
            "Scheduled Date",
            "Drain Balance",
            "Applied Date",
            "Pay Date",
            "Active Goal",
            "Progress",
            "Chart Label",
        }
        for row in sheet.iter_rows():
            header = self._section_header_for_cell(sheet, row[0].row)
            for cell in row:
                if hasattr(cell.value, "year"):
                    cell.number_format = "mmm d, yyyy"
                elif (
                    isinstance(cell.value, (int, float))
                    and not isinstance(cell.value, bool)
                    and header.get(cell.column) not in non_currency_headers
                ):
                    cell.number_format = "$#,##0.00"

        progress_col = self._header_column(sheet, "Progress", detail_header_row)
        for cell in sheet[get_column_letter(progress_col)][detail_header_row:]:
            cell.number_format = "0.0%"

        label_col = self._header_column(sheet, "Chart Label", detail_header_row)
        sheet.column_dimensions[get_column_letter(label_col)].hidden = True
        for column_index, column in enumerate(sheet.columns, start=1):
            max_length = max(
                len(str(cell.value)) if cell.value is not None else 0 for cell in column
            )
            sheet.column_dimensions[get_column_letter(column_index)].width = min(
                max(max_length + 2, 12),
                32,
            )
        sheet.freeze_panes = f"A{detail_header_row + 1}"
        sheet.auto_filter.ref = f"A{detail_header_row}:{get_column_letter(sheet.max_column)}{sheet.max_row}"

    def _section_header_for_cell(self, sheet: Worksheet, row_index: int) -> dict[int, str]:
        """Return the nearest header row mapping above a cell row."""
        header_map: dict[int, str] = {}
        for row in range(row_index, 0, -1):
            values = [cell.value for cell in sheet[row]]
            if "Goal Name" in values or "Pay Date" in values or "Scheduled Date" in values:
                return {
                    cell.column: str(cell.value)
                    for cell in sheet[row]
                    if cell.value is not None
                }
        return header_map

    def _write_dashboard(
        self,
        sheet: Worksheet,
        summaries: list[PayPeriodSummary],
        forecast: ForecastSummary | None = None,
        scenario_comparison: ScenarioComparison | None = None,
        target_result: DebtFreeTargetResult | None = None,
    ) -> None:
        sheet["A1"] = "DebtSnowball Dashboard"
        sheet["A1"].font = Font(bold=True, size=16)
        sheet.merge_cells("A1:D1")

        sheet["A3"] = "Metric"
        sheet["B3"] = "Value"

        metrics = self._dashboard_metrics(
            summaries,
            forecast,
            scenario_comparison,
            target_result,
        )
        for index, (label, value, number_format) in enumerate(metrics, start=4):
            sheet.cell(row=index, column=1, value=label)
            value_cell = sheet.cell(row=index, column=2, value=value)
            value_cell.number_format = number_format

        self._format_table(sheet, currency_columns=[], date_columns=[], header_row=3)
        sheet.auto_filter.ref = f"A3:B{len(metrics) + 3}"
        sheet.freeze_panes = "A4"
        sheet.column_dimensions["A"].width = 28
        sheet.column_dimensions["B"].width = 18
        sheet["A1"].alignment = Alignment(horizontal="left")

        self._add_savings_chart(sheet, len(summaries))
        self._add_debt_chart(sheet, len(summaries))

    def _dashboard_metrics(
        self,
        summaries: list[PayPeriodSummary],
        forecast: ForecastSummary | None = None,
        scenario_comparison: ScenarioComparison | None = None,
        target_result: DebtFreeTargetResult | None = None,
    ) -> list[tuple[str, object, str]]:
        if not summaries:
            metrics: list[tuple[str, object, str]] = []
        else:
            first_summary = summaries[0]
            last_summary = summaries[-1]
            starting_savings = money(
                first_summary.savings_balance - first_summary.savings_contribution,
            )
            current_savings = last_summary.savings_balance
            savings_goal = last_summary.savings_goal
            remaining_to_goal = max(savings_goal - current_savings, ZERO_MONEY)
            current_debt = self._active_debt_total(last_summary)
            paid_off_count = len(last_summary.paid_off_debts)
            total_snowball = money(
                sum(
                    (summary.snowball_payment for summary in summaries),
                    ZERO_MONEY,
                )
            )
            total_minimums = money(
                sum((summary.debt_minimums for summary in summaries), ZERO_MONEY)
            )

            metrics = [
                ("Starting Savings", starting_savings, "$#,##0.00"),
                ("Current Savings", current_savings, "$#,##0.00"),
                ("Savings Goal", savings_goal, "$#,##0.00"),
                ("Remaining To Goal", remaining_to_goal, "$#,##0.00"),
                ("Current Active Debt", current_debt, "$#,##0.00"),
                ("Paid-Off Debts", paid_off_count, "0"),
                ("Total Snowball Paid", total_snowball, "$#,##0.00"),
                ("Total Minimums Paid", total_minimums, "$#,##0.00"),
            ]
            if last_summary.active_savings_goal_name is not None:
                active_stage_result = self._active_savings_stage_result(
                    forecast,
                    last_summary.active_savings_goal_name,
                )
                metrics.extend(
                    [
                        ("Active Savings Goal", last_summary.active_savings_goal_name, "@"),
                        (
                            "Active Savings Target",
                            last_summary.active_savings_target,
                            "$#,##0.00",
                        ),
                        (
                            "Active Savings Target Date",
                            last_summary.savings_goal_target_date,
                            "mmm d, yyyy",
                        ),
                        (
                            "Current Savings Progress",
                            last_summary.goal_progress_percentage,
                            "0.0%",
                        ),
                        (
                            "Projected Balance At Deadline",
                            self._cell_value(
                                active_stage_result.projected_balance_at_deadline
                                if active_stage_result is not None
                                else None
                            ),
                            "$#,##0.00",
                        ),
                        (
                            "Projected Shortfall",
                            self._cell_value(
                                active_stage_result.projected_shortfall
                                if active_stage_result is not None
                                else None
                            ),
                            "$#,##0.00",
                        ),
                        (
                            "Goal Feasible",
                            self._yes_no(
                                active_stage_result.feasible_under_current_plan
                                if active_stage_result is not None
                                else None
                            ),
                            "@",
                        ),
                        (
                            "Snowball Currently Reduced",
                            "Yes" if last_summary.snowball_reduction > 0 else "No",
                            "@",
                        ),
                    ]
                )

        if forecast is not None:
            metrics.extend(
                [
                    (
                        "Estimated Debt-Free Date",
                        forecast.debt_free_date,
                        "mmm d, yyyy",
                    ),
                    (
                        "Estimated Savings Goal Date",
                        forecast.savings_goal_date,
                        "mmm d, yyyy",
                    ),
                    (
                        "Remaining Forecast Interest",
                        self._cell_value(forecast.total_interest_paid),
                        "$#,##0.00",
                    ),
                ]
            )

        metrics.extend(self._scenario_dashboard_metrics(scenario_comparison))
        metrics.extend(self._target_dashboard_metrics(target_result))
        return metrics

    def _active_savings_stage_result(
        self,
        forecast: ForecastSummary | None,
        goal_name: str,
    ):
        if forecast is None:
            return None

        return next(
            (
                stage
                for stage in forecast.savings_stage_results
                if stage.goal_name == goal_name
            ),
            None,
        )

    def _target_dashboard_metrics(
        self,
        target_result: DebtFreeTargetResult | None,
    ) -> list[tuple[str, object, str]]:
        if target_result is None:
            return []

        required_payment = (
            None
            if target_result.required_extra_per_paycheck is None
            else self._cell_value(target_result.required_extra_per_paycheck)
        )
        status = (
            "Target not reachable"
            if target_result.calculation_status == DebtFreeTargetStatus.UNREACHABLE
            else "Target met"
            if target_result.target_met
            else "Not configured"
        )

        return [
            ("Target Date", target_result.target_date, "mmm d, yyyy"),
            ("Required Extra Per Paycheck", required_payment, "$#,##0.00"),
            (
                "Target Projected Debt-Free Date",
                target_result.projected_debt_free_date,
                "mmm d, yyyy",
            ),
            ("Target Status", status, "@"),
        ]

    def _scenario_dashboard_metrics(
        self,
        scenario_comparison: ScenarioComparison | None,
    ) -> list[tuple[str, object, str]]:
        best = self._best_completed_scenario(scenario_comparison)
        if best is None:
            return [
                ("Best Scenario", "No scenario completed within horizon", "@"),
                ("Earliest Debt-Free Date", None, "mmm d, yyyy"),
                ("Maximum Interest Saved", self._cell_value(ZERO_MONEY), "$#,##0.00"),
                ("Extra Payment Required", self._cell_value(ZERO_MONEY), "$#,##0.00"),
            ]

        baseline = scenario_comparison.baseline.forecast
        return [
            ("Best Scenario", best.name, "@"),
            ("Earliest Debt-Free Date", best.forecast.debt_free_date, "mmm d, yyyy"),
            (
                "Maximum Interest Saved",
                self._cell_value(
                    baseline.total_interest_paid - best.forecast.total_interest_paid
                ),
                "$#,##0.00",
            ),
            (
                "Extra Payment Required",
                self._cell_value(best.extra_per_paycheck),
                "$#,##0.00",
            ),
        ]

    def _write_debt_free_target(
        self,
        sheet: Worksheet,
        result: DebtFreeTargetResult,
    ) -> None:
        """Write debt-free target calculator output."""
        sheet["A1"] = "Debt-Free Target"
        sheet["A1"].font = Font(bold=True, size=16)
        sheet.merge_cells("A1:B1")

        sheet["A3"] = "Metric"
        sheet["B3"] = "Value"
        rows = self._debt_free_target_rows(result)
        for row_index, (label, value, number_format) in enumerate(rows, start=4):
            sheet.cell(row=row_index, column=1, value=label)
            value_cell = sheet.cell(row=row_index, column=2, value=value)
            value_cell.number_format = number_format

        self._format_table(sheet, currency_columns=[], date_columns=[], header_row=3)
        sheet.auto_filter.ref = f"A3:B{len(rows) + 3}"
        sheet.freeze_panes = "A4"
        sheet.column_dimensions["A"].width = 34
        sheet.column_dimensions["B"].width = 32

    def _debt_free_target_rows(
        self,
        result: DebtFreeTargetResult,
    ) -> list[tuple[str, object, str]]:
        required_payment = (
            None
            if result.required_extra_per_paycheck is None
            else self._cell_value(result.required_extra_per_paycheck)
        )
        interest_difference = (
            result.baseline_total_interest - result.total_interest
        )
        days_accelerated = self._days_saved(
            result.baseline_debt_free_date,
            result.projected_debt_free_date,
        )
        status_message = result.message or result.calculation_status.value

        return [
            ("Target Date", result.target_date, "mmm d, yyyy"),
            ("Required Extra Per Paycheck", required_payment, "$#,##0.00"),
            (
                "Projected Debt-Free Date",
                result.projected_debt_free_date,
                "mmm d, yyyy",
            ),
            ("Target Met", result.target_met, "@"),
            ("Baseline Debt-Free Date", result.baseline_debt_free_date, "mmm d, yyyy"),
            ("Days Accelerated", days_accelerated, "0"),
            (
                "Maximum Extra Tested",
                self._cell_value(result.maximum_extra_tested),
                "$#,##0.00",
            ),
            ("Search Precision", self._cell_value(result.precision), "$#,##0.00"),
            ("Iterations Used", result.iterations_used, "0"),
            (
                "Total Interest at Required Payment",
                self._cell_value(result.total_interest),
                "$#,##0.00",
            ),
            (
                "Interest Difference vs Baseline",
                self._cell_value(interest_difference),
                "$#,##0.00",
            ),
            (
                "Total Snowball Paid",
                self._cell_value(result.total_snowball_paid),
                "$#,##0.00",
            ),
            ("Ending Debt", self._cell_value(result.ending_debt), "$#,##0.00"),
            ("Status Message", status_message, "@"),
        ]

    def _write_scenario_comparison(
        self,
        sheet: Worksheet,
        scenario_comparison: ScenarioComparison | None,
    ) -> None:
        """Write scenario summary, deltas, payoff comparison, and charts."""
        sheet["A1"] = "Scenario Comparison"
        sheet["A1"].font = Font(bold=True, size=16)
        sheet.merge_cells("A1:M1")

        summary_header_row = 3
        summary_rows = self._scenario_summary_rows(scenario_comparison)
        summary_headers = [
            "Scenario",
            "Extra Per Paycheck",
            "Debt-Free Date",
            "Days Saved",
            "Savings Goal Date",
            "Savings Goal Days Changed",
            "Total Interest",
            "Interest Saved",
            "Total Snowball Paid",
            "Additional Snowball Paid",
            "Ending Debt",
            "Ending Savings",
            "Completed",
        ]
        self._write_header(sheet, summary_header_row, summary_headers)
        for row_index, row in enumerate(summary_rows, start=summary_header_row + 1):
            for column_index, value in enumerate(row, start=1):
                sheet.cell(row=row_index, column=column_index, value=value)

        deltas_header_row = summary_header_row + max(len(summary_rows), 1) + 3
        sheet.cell(row=deltas_header_row - 1, column=1, value="Scenario Deltas")
        sheet.cell(row=deltas_header_row - 1, column=1).font = Font(bold=True)
        delta_headers = [
            "Scenario",
            "Debt-Free Days Saved",
            "Savings Goal Days Changed",
            "Interest Saved",
            "Additional Snowball Paid",
            "Ending Debt Difference",
            "Ending Savings Difference",
        ]
        self._write_header(sheet, deltas_header_row, delta_headers)
        delta_rows = self._scenario_delta_rows(scenario_comparison)
        for row_index, row in enumerate(delta_rows, start=deltas_header_row + 1):
            for column_index, value in enumerate(row, start=1):
                sheet.cell(row=row_index, column=column_index, value=value)

        payoff_header_row = deltas_header_row + max(len(delta_rows), 1) + 3
        sheet.cell(row=payoff_header_row - 1, column=1, value="Debt Payoff Comparison")
        sheet.cell(row=payoff_header_row - 1, column=1).font = Font(bold=True)
        payoff_headers = self._debt_payoff_headers(scenario_comparison)
        self._write_header(sheet, payoff_header_row, payoff_headers)
        payoff_rows = self._debt_payoff_rows(scenario_comparison)
        for row_index, row in enumerate(payoff_rows, start=payoff_header_row + 1):
            for column_index, value in enumerate(row, start=1):
                sheet.cell(row=row_index, column=column_index, value=value)

        chart_data_header_row = payoff_header_row + max(len(payoff_rows), 1) + 3
        self._write_scenario_chart_data(
            sheet,
            scenario_comparison,
            chart_data_header_row,
        )
        self._format_scenario_sheet(
            sheet=sheet,
            summary_header_row=summary_header_row,
            summary_row_count=len(summary_rows),
            deltas_header_row=deltas_header_row,
            delta_row_count=len(delta_rows),
            payoff_header_row=payoff_header_row,
            payoff_column_count=len(payoff_headers),
            chart_data_header_row=chart_data_header_row,
        )
        self._add_scenario_charts(sheet, chart_data_header_row, len(summary_rows))

    def _scenario_summary_rows(
        self,
        scenario_comparison: ScenarioComparison | None,
    ) -> list[list[object]]:
        if scenario_comparison is None:
            return []

        baseline = scenario_comparison.baseline.forecast
        rows = [
            self._scenario_summary_row(
                scenario_comparison.baseline,
                baseline,
                is_baseline=True,
            )
        ]
        rows.extend(
            self._scenario_summary_row(scenario, baseline)
            for scenario in scenario_comparison.scenarios
        )
        return rows

    def _scenario_summary_row(
        self,
        scenario: ScenarioResult,
        baseline,
        is_baseline: bool = False,
    ) -> list[object]:
        forecast = scenario.forecast
        debt_free_days_saved = self._days_saved(
            baseline.debt_free_date,
            forecast.debt_free_date,
        )
        savings_goal_days_changed = self._days_saved(
            baseline.savings_goal_date,
            forecast.savings_goal_date,
        )

        return [
            scenario.name,
            self._cell_value(scenario.extra_per_paycheck),
            forecast.debt_free_date,
            0
            if is_baseline
            else debt_free_days_saved
            if debt_free_days_saved is not None
            else None,
            forecast.savings_goal_date,
            0
            if is_baseline
            else savings_goal_days_changed
            if savings_goal_days_changed is not None
            else None,
            self._cell_value(forecast.total_interest_paid),
            self._cell_value(ZERO_MONEY)
            if is_baseline
            else self._cell_value(
                baseline.total_interest_paid - forecast.total_interest_paid
            ),
            self._cell_value(forecast.total_snowball_payments),
            self._cell_value(ZERO_MONEY)
            if is_baseline
            else self._cell_value(
                forecast.total_snowball_payments - baseline.total_snowball_payments
            ),
            self._cell_value(forecast.remaining_debt),
            self._cell_value(forecast.ending_savings),
            forecast.completed,
        ]

    def _scenario_delta_rows(
        self,
        scenario_comparison: ScenarioComparison | None,
    ) -> list[list[object]]:
        if scenario_comparison is None:
            return []

        baseline = scenario_comparison.baseline.forecast
        return [
            [
                scenario.name,
                self._days_saved(
                    baseline.debt_free_date,
                    scenario.forecast.debt_free_date,
                ),
                self._days_saved(
                    baseline.savings_goal_date,
                    scenario.forecast.savings_goal_date,
                ),
                self._cell_value(
                    baseline.total_interest_paid - scenario.forecast.total_interest_paid
                ),
                self._cell_value(
                    scenario.forecast.total_snowball_payments
                    - baseline.total_snowball_payments
                ),
                self._cell_value(
                    scenario.forecast.remaining_debt - baseline.remaining_debt
                ),
                self._cell_value(
                    scenario.forecast.ending_savings - baseline.ending_savings
                ),
            ]
            for scenario in scenario_comparison.scenarios
        ]

    def _debt_payoff_headers(
        self,
        scenario_comparison: ScenarioComparison | None,
    ) -> list[str]:
        if scenario_comparison is None:
            return ["Debt", "Baseline"]

        return [
            "Debt",
            *[scenario.name for scenario in self._scenario_results(scenario_comparison)],
        ]

    def _debt_payoff_rows(
        self,
        scenario_comparison: ScenarioComparison | None,
    ) -> list[list[object]]:
        if scenario_comparison is None:
            return []

        results = self._scenario_results(scenario_comparison)
        debt_names = self._scenario_debt_names(results)
        return [
            [
                debt_name,
                *[
                    self._payoff_date_for_debt(result, debt_name)
                    for result in results
                ],
            ]
            for debt_name in debt_names
        ]

    def _write_scenario_chart_data(
        self,
        sheet: Worksheet,
        scenario_comparison: ScenarioComparison | None,
        header_row: int,
    ) -> None:
        self._write_header(
            sheet,
            header_row,
            ["Scenario", "Total Interest", "Days To Debt-Free"],
        )
        if scenario_comparison is None:
            return

        for row_index, result in enumerate(
            self._scenario_results(scenario_comparison),
            start=header_row + 1,
        ):
            forecast = result.forecast
            days_to_debt_free = self._days_between(
                forecast.forecast_start_date,
                forecast.debt_free_date,
            )
            sheet.cell(row=row_index, column=1, value=result.name)
            sheet.cell(
                row=row_index,
                column=2,
                value=self._cell_value(forecast.total_interest_paid),
            )
            sheet.cell(row=row_index, column=3, value=days_to_debt_free)

    def _format_scenario_sheet(
        self,
        sheet: Worksheet,
        summary_header_row: int,
        summary_row_count: int,
        deltas_header_row: int,
        delta_row_count: int,
        payoff_header_row: int,
        payoff_column_count: int,
        chart_data_header_row: int,
    ) -> None:
        header_rows = [
            summary_header_row,
            deltas_header_row,
            payoff_header_row,
            chart_data_header_row,
        ]
        for header_row in header_rows:
            for cell in sheet[header_row]:
                if cell.value is not None:
                    self._style_header_cell(cell)

        currency_columns = ["B", "G", "H", "I", "J", "K", "L"]
        date_columns = ["C", "E"]
        integer_columns = ["D", "F"]
        for column_letter in currency_columns:
            for cell in sheet[column_letter][summary_header_row:]:
                cell.number_format = "$#,##0.00"
        for column_letter in date_columns:
            for cell in sheet[column_letter][summary_header_row:]:
                cell.number_format = "mmm d, yyyy"
        for column_letter in integer_columns:
            for cell in sheet[column_letter][summary_header_row:]:
                cell.number_format = "0"

        for column_letter in ["D", "E", "F", "G"]:
            for cell in sheet[column_letter][deltas_header_row:payoff_header_row - 1]:
                cell.number_format = "$#,##0.00"
        for column_letter in ["B", "C"]:
            for cell in sheet[column_letter][deltas_header_row:payoff_header_row - 1]:
                cell.number_format = "0"

        for row in sheet.iter_rows(
            min_row=payoff_header_row + 1,
            max_row=sheet.max_row,
            min_col=2,
            max_col=max(payoff_column_count, 2),
        ):
            for cell in row:
                if hasattr(cell.value, "year"):
                    cell.number_format = "mmm d, yyyy"

        sheet.column_dimensions["A"].width = 26
        for column_index, column in enumerate(sheet.columns, start=1):
            max_length = max(
                len(str(cell.value)) if cell.value is not None else 0 for cell in column
            )
            sheet.column_dimensions[get_column_letter(column_index)].width = min(
                max(max_length + 2, 12),
                34,
            )

        sheet.freeze_panes = "A4"
        sheet.auto_filter.ref = (
            f"A{summary_header_row}:M{summary_header_row + summary_row_count}"
        )
        sheet.cell(row=chart_data_header_row - 1, column=1, value="Chart Data")
        sheet.cell(row=chart_data_header_row - 1, column=1).font = Font(bold=True)

    def _add_scenario_charts(
        self,
        sheet: Worksheet,
        chart_data_header_row: int,
        scenario_count: int,
    ) -> None:
        if scenario_count == 0:
            return

        max_row = chart_data_header_row + scenario_count
        categories = Reference(
            sheet,
            min_col=1,
            min_row=chart_data_header_row + 1,
            max_row=max_row,
        )

        interest_chart = BarChart()
        interest_chart.title = "Interest by Scenario"
        interest_chart.y_axis.title = "Total Interest"
        interest_chart.x_axis.title = "Scenario"
        interest_chart.add_data(
            Reference(sheet, min_col=2, min_row=chart_data_header_row, max_row=max_row),
            titles_from_data=True,
        )
        interest_chart.set_categories(categories)
        interest_chart.height = 7
        interest_chart.width = 14
        sheet.add_chart(interest_chart, "O3")

        timeline_chart = BarChart()
        timeline_chart.title = "Debt-Free Timeline by Scenario"
        timeline_chart.y_axis.title = "Days To Debt-Free"
        timeline_chart.x_axis.title = "Scenario"
        timeline_chart.add_data(
            Reference(sheet, min_col=3, min_row=chart_data_header_row, max_row=max_row),
            titles_from_data=True,
        )
        timeline_chart.set_categories(categories)
        timeline_chart.height = 7
        timeline_chart.width = 14
        sheet.add_chart(timeline_chart, "O20")

    def _write_header(
        self,
        sheet: Worksheet,
        row: int,
        headers: list[str],
    ) -> None:
        for column_index, header in enumerate(headers, start=1):
            cell = sheet.cell(row=row, column=column_index, value=header)
            self._style_header_cell(cell)

    def _style_header_cell(self, cell) -> None:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True)

    def _scenario_results(
        self,
        scenario_comparison: ScenarioComparison,
    ) -> list[ScenarioResult]:
        return [scenario_comparison.baseline, *scenario_comparison.scenarios]

    def _scenario_debt_names(self, results: list[ScenarioResult]) -> list[str]:
        debt_names = []
        seen = set()
        for result in results:
            for payoff in result.debt_payoffs:
                if payoff.debt_name in seen:
                    continue

                seen.add(payoff.debt_name)
                debt_names.append(payoff.debt_name)

        return debt_names

    def _payoff_date_for_debt(
        self,
        result: ScenarioResult,
        debt_name: str,
    ) -> object:
        for payoff in result.debt_payoffs:
            if payoff.debt_name != debt_name:
                continue

            return payoff.payoff_date or "Not paid within horizon"

        return "Not paid within horizon"

    def _days_saved(
        self,
        baseline_date,
        scenario_date,
    ) -> int | None:
        if baseline_date is None or scenario_date is None:
            return None

        return (baseline_date - scenario_date).days

    def _days_between(
        self,
        start_date,
        end_date,
    ) -> int | None:
        if start_date is None or end_date is None:
            return None

        return (end_date - start_date).days

    def _best_completed_scenario(
        self,
        scenario_comparison: ScenarioComparison | None,
    ) -> ScenarioResult | None:
        if scenario_comparison is None:
            return None

        completed = [
            result
            for result in self._scenario_results(scenario_comparison)
            if result.forecast.debt_free_date is not None
        ]
        if not completed:
            return None

        return min(
            completed,
            key=lambda result: (
                result.forecast.debt_free_date,
                result.extra_per_paycheck,
            ),
        )

    def _write_forecast(
        self,
        sheet: Worksheet,
        forecast: ForecastSummary | None,
    ) -> None:
        sheet["A1"] = "Forecast"
        sheet["A1"].font = Font(bold=True, size=16)
        sheet.merge_cells("A1:E1")

        summary_header_row = 3
        sheet.cell(row=summary_header_row, column=1, value="Summary")
        sheet.cell(row=summary_header_row, column=1).font = Font(bold=True)

        summary_rows = self._forecast_summary_rows(forecast)
        for row_index, (label, value, number_format) in enumerate(
            summary_rows,
            start=summary_header_row + 1,
        ):
            sheet.cell(row=row_index, column=1, value=label)
            value_cell = sheet.cell(row=row_index, column=2, value=value)
            value_cell.number_format = number_format

        debt_header_row = summary_header_row + len(summary_rows) + 3
        sheet.cell(row=debt_header_row, column=1, value="Debt Forecast")
        sheet.cell(row=debt_header_row, column=1).font = Font(bold=True)
        debt_table_header_row = debt_header_row + 1
        debt_headers = [
            "Debt",
            "Starting Balance",
            "Estimated Payoff Date",
            "Interest Paid",
            "Total Paid",
        ]
        for column_index, header in enumerate(debt_headers, start=1):
            sheet.cell(row=debt_table_header_row, column=column_index, value=header)

        debt_rows = forecast.debt_payoffs if forecast is not None else []
        for row_index, payoff in enumerate(debt_rows, start=debt_table_header_row + 1):
            sheet.cell(row=row_index, column=1, value=payoff.debt_name)
            sheet.cell(row=row_index, column=2, value=self._cell_value(payoff.starting_balance))
            sheet.cell(row=row_index, column=3, value=payoff.payoff_date)
            sheet.cell(row=row_index, column=4, value=self._cell_value(payoff.total_interest_paid))
            sheet.cell(row=row_index, column=5, value=self._cell_value(payoff.total_paid))

        timeline_header_row = debt_table_header_row + max(len(debt_rows), 1) + 3
        sheet.cell(row=timeline_header_row, column=1, value="Forecast Timeline")
        sheet.cell(row=timeline_header_row, column=1).font = Font(bold=True)
        timeline_table_header_row = timeline_header_row + 1
        timeline_headers = [
            "Paycheck Date",
            "Debt Balance",
            "Savings Balance",
            "Interest Paid",
            "Snowball Paid",
            "Chart Label",
        ]
        for column_index, header in enumerate(timeline_headers, start=1):
            sheet.cell(row=timeline_table_header_row, column=column_index, value=header)

        timeline_rows = forecast.periods if forecast is not None else []
        for row_index, period in enumerate(
            timeline_rows,
            start=timeline_table_header_row + 1,
        ):
            sheet.cell(row=row_index, column=1, value=period.paycheck_date)
            sheet.cell(row=row_index, column=2, value=self._cell_value(period.total_debt_balance))
            sheet.cell(row=row_index, column=3, value=self._cell_value(period.savings_balance))
            sheet.cell(row=row_index, column=4, value=self._cell_value(period.interest_paid))
            sheet.cell(row=row_index, column=5, value=self._cell_value(period.snowball_paid))
            sheet.cell(
                row=row_index,
                column=6,
                value=self._chart_date_label(period.paycheck_date, len(timeline_rows)),
            )

        self._format_forecast_sheet(
            sheet=sheet,
            debt_table_header_row=debt_table_header_row,
            timeline_table_header_row=timeline_table_header_row,
        )
        self._add_forecast_charts(sheet, timeline_table_header_row, len(timeline_rows))

    def _forecast_summary_rows(
        self,
        forecast: ForecastSummary | None,
    ) -> list[tuple[str, object, str]]:
        if forecast is None:
            return [
                ("Estimated Debt-Free Date", None, "mmm d, yyyy"),
                ("Estimated Savings Goal Date", None, "mmm d, yyyy"),
                ("Remaining Debt", self._cell_value(ZERO_MONEY), "$#,##0.00"),
                ("Ending Savings", self._cell_value(ZERO_MONEY), "$#,##0.00"),
                ("Total Interest Remaining", self._cell_value(ZERO_MONEY), "$#,##0.00"),
                (
                    "Total Future Minimum Payments",
                    self._cell_value(ZERO_MONEY),
                    "$#,##0.00",
                ),
                (
                    "Total Future Snowball Payments",
                    self._cell_value(ZERO_MONEY),
                    "$#,##0.00",
                ),
            ]

        return [
            ("Estimated Debt-Free Date", forecast.debt_free_date, "mmm d, yyyy"),
            ("Estimated Savings Goal Date", forecast.savings_goal_date, "mmm d, yyyy"),
            ("Remaining Debt", self._cell_value(forecast.remaining_debt), "$#,##0.00"),
            ("Ending Savings", self._cell_value(forecast.ending_savings), "$#,##0.00"),
            (
                "Total Interest Remaining",
                self._cell_value(forecast.total_interest_paid),
                "$#,##0.00",
            ),
            (
                "Total Future Minimum Payments",
                self._cell_value(forecast.total_minimum_payments),
                "$#,##0.00",
            ),
            (
                "Total Future Snowball Payments",
                self._cell_value(forecast.total_snowball_payments),
                "$#,##0.00",
            ),
        ]

    def _format_forecast_sheet(
        self,
        sheet: Worksheet,
        debt_table_header_row: int,
        timeline_table_header_row: int,
    ) -> None:
        header_fill = PatternFill("solid", fgColor="1F4E78")
        header_font = Font(color="FFFFFF", bold=True)

        for header_row in [debt_table_header_row, timeline_table_header_row]:
            for cell in sheet[header_row]:
                cell.fill = header_fill
                cell.font = header_font

        for cell in sheet["A"]:
            if cell.row > 1:
                cell.number_format = "mmm d, yyyy"

        for column_letter in ["B", "D", "E"]:
            for cell in sheet[column_letter][1:]:
                cell.number_format = "$#,##0.00"

        for cell in sheet["C"][1:]:
            if isinstance(cell.value, (int, float)):
                cell.number_format = "$#,##0.00"
            else:
                cell.number_format = "mmm d, yyyy"

        for column_index, column in enumerate(sheet.columns, start=1):
            max_length = max(
                len(str(cell.value)) if cell.value is not None else 0 for cell in column
            )
            sheet.column_dimensions[get_column_letter(column_index)].width = min(
                max(max_length + 2, 12),
                32,
            )

        sheet.freeze_panes = "A4"
        sheet.auto_filter.ref = f"A{timeline_table_header_row}:E{sheet.max_row}"

    def _add_forecast_charts(
        self,
        sheet: Worksheet,
        timeline_table_header_row: int,
        timeline_count: int,
    ) -> None:
        if timeline_count == 0:
            return

        max_row = timeline_table_header_row + timeline_count
        categories = Reference(
            sheet,
            min_col=6,
            min_row=timeline_table_header_row + 1,
            max_row=max_row,
        )

        debt_chart = LineChart()
        debt_chart.title = "Total Debt Balance Over Time"
        debt_chart.y_axis.title = "Debt Balance"
        debt_chart.x_axis.title = "Paycheck Date"
        debt_data = Reference(
            sheet,
            min_col=2,
            min_row=timeline_table_header_row,
            max_row=max_row,
        )
        debt_chart.add_data(debt_data, titles_from_data=True)
        debt_chart.set_categories(categories)
        self._format_date_chart_axis(debt_chart, timeline_count)
        debt_chart.height = 7
        debt_chart.width = 14
        sheet.add_chart(debt_chart, "G3")

        savings_chart = LineChart()
        savings_chart.title = "Savings Balance Over Time"
        savings_chart.y_axis.title = "Savings Balance"
        savings_chart.x_axis.title = "Paycheck Date"
        savings_data = Reference(
            sheet,
            min_col=3,
            min_row=timeline_table_header_row,
            max_row=max_row,
        )
        savings_chart.add_data(savings_data, titles_from_data=True)
        savings_chart.set_categories(categories)
        self._format_date_chart_axis(savings_chart, timeline_count)
        savings_chart.height = 7
        savings_chart.width = 14
        sheet.add_chart(savings_chart, "G20")
        sheet.column_dimensions["F"].hidden = True

    def _add_savings_chart(self, sheet: Worksheet, summary_count: int) -> None:
        if summary_count == 0:
            return

        savings_sheet = sheet.parent["Savings Progress"]
        header_row, balance_col, label_col = self._savings_chart_columns(savings_sheet)

        chart = LineChart()
        chart.title = "Savings Growth"
        chart.y_axis.title = "Savings Balance"
        chart.x_axis.title = "Pay Date"
        chart.style = 13

        data = Reference(
            savings_sheet,
            min_col=balance_col,
            min_row=header_row,
            max_row=header_row + summary_count,
        )
        categories = Reference(
            savings_sheet,
            min_col=label_col,
            min_row=header_row + 1,
            max_row=header_row + summary_count,
        )
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(categories)
        self._format_date_chart_axis(chart, summary_count)
        chart.height = 7
        chart.width = 14

        sheet.add_chart(chart, "D3")
        savings_sheet.column_dimensions[get_column_letter(label_col)].hidden = True

    def _savings_chart_columns(self, sheet: Worksheet) -> tuple[int, int, int]:
        """Return header row, balance column, and label column for savings charts."""
        for row in sheet.iter_rows():
            values = [cell.value for cell in row]
            if "Ending Savings Balance" in values and "Chart Label" in values:
                header_row = row[0].row
                return (
                    header_row,
                    values.index("Ending Savings Balance") + 1,
                    values.index("Chart Label") + 1,
                )
            if "Savings Balance" in values and "Chart Label" in values:
                header_row = row[0].row
                return (
                    header_row,
                    values.index("Savings Balance") + 1,
                    values.index("Chart Label") + 1,
                )

        return 1, 3, 6

    def _add_debt_chart(self, sheet: Worksheet, summary_count: int) -> None:
        if summary_count == 0:
            return

        chart = LineChart()
        chart.title = "Debt Reduction"
        chart.y_axis.title = "Active Debt Total"
        chart.x_axis.title = "Pay Date"
        chart.style = 12

        data = Reference(
            sheet.parent["Pay Period Summaries"],
            min_col=self._header_column(sheet.parent["Pay Period Summaries"], "Active Debt Total"),
            min_row=1,
            max_row=summary_count + 1,
        )
        categories = Reference(
            sheet.parent["Pay Period Summaries"],
            min_col=self._header_column(sheet.parent["Pay Period Summaries"], "Chart Label"),
            min_row=2,
            max_row=summary_count + 1,
        )
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(categories)
        self._format_date_chart_axis(chart, summary_count)
        chart.height = 7
        chart.width = 14

        sheet.add_chart(chart, "D20")
        label_col = self._header_column(sheet.parent["Pay Period Summaries"], "Chart Label")
        sheet.parent["Pay Period Summaries"].column_dimensions[
            get_column_letter(label_col)
        ].hidden = True

    def _header_column(self, sheet: Worksheet, header: str, header_row: int = 1) -> int:
        """Return the 1-based column for a header."""
        for cell in sheet[header_row]:
            if cell.value == header:
                return cell.column

        raise ValueError(f"Missing expected header: {header}")

    def _format_date_chart_axis(self, chart: LineChart, point_count: int) -> None:
        """Format date-based chart categories with readable labels."""
        chart.x_axis.number_format = "@"
        chart.x_axis.majorTickMark = "out"
        chart.x_axis.tickLblPos = "low"
        chart.x_axis.tickLblSkip = max(1, (point_count + 7) // 8)
        chart.x_axis.tickMarkSkip = max(1, (point_count + 7) // 8)

    def _chart_date_label(self, value, point_count: int) -> str:
        """Return a text date label for chart categories."""
        return value.strftime("%b %Y" if point_count > 24 else "%b %d")

    def _format_table(
        self,
        sheet: Worksheet,
        currency_columns: list[int] | None = None,
        date_columns: list[int] | None = None,
        header_row: int = 1,
    ) -> None:
        currency_columns = currency_columns or []
        date_columns = date_columns or []

        header_fill = PatternFill("solid", fgColor="1F4E78")
        header_font = Font(color="FFFFFF", bold=True)

        for cell in sheet[header_row]:
            cell.fill = header_fill
            cell.font = header_font

        if sheet.max_row >= 1 and sheet.max_column >= 1:
            sheet.auto_filter.ref = sheet.dimensions

        for column_index in currency_columns:
            for cell in sheet[get_column_letter(column_index)][1:]:
                cell.number_format = "$#,##0.00"

        for column_index in date_columns:
            for cell in sheet[get_column_letter(column_index)][1:]:
                cell.number_format = "mmm d, yyyy"

        for column_index, column in enumerate(sheet.columns, start=1):
            max_length = max(
                len(str(cell.value)) if cell.value is not None else 0 for cell in column
            )
            sheet.column_dimensions[get_column_letter(column_index)].width = min(
                max(max_length + 2, 12),
                28,
            )

        sheet.freeze_panes = "A2"

    def _active_debt_total(self, summary: PayPeriodSummary):
        return money(
            sum((debt.balance for debt in summary.active_debt_balances), ZERO_MONEY)
        )

    def _append_row(self, sheet: Worksheet, values: list[object]) -> None:
        """Append a row while converting Decimal money at the Excel boundary."""
        sheet.append([self._cell_value(value) for value in values])

    def _cell_value(self, value: object) -> object:
        if hasattr(value, "quantize"):
            return excel_number(value)

        return value

    def _money_cell(self, value: object) -> object:
        if value is None:
            return None
        return excel_number(money(value))

    def _date_cell(self, value: object) -> object:
        if value is None or hasattr(value, "year"):
            return value
        return date.fromisoformat(str(value))

    def _yes_no(self, value: bool | None) -> str | None:
        if value is None:
            return None

        return "Yes" if value else "No"
