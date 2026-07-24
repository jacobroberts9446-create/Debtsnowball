import json
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from app.config import Config
from app.excel_writer import ExcelWriter
from app.models import DebtFreeTargetStatus
from app.workflows.workbook_export import (
    build_debt_free_target_result,
    build_scenario_comparison,
)


def config_data(scenarios_marker="missing", target_marker="missing"):
    data = {
        "budget": {
            "paycheck": 500,
            "first_paycheck": "2026-01-02",
            "rent_per_paycheck": 0,
            "insurance_per_paycheck": 0,
            "personal_per_paycheck": 0,
            "starting_savings": 0,
            "savings_goal": 0,
            "snowball_split": 0.5,
        },
        "bills": [],
        "debts": [
            {
                "name": "Debt A",
                "balance": 500,
                "apr": 0,
                "minimum": 50,
                "due_day": 10,
                "snowball_order": 1,
            },
            {
                "name": "Debt B",
                "balance": 700,
                "apr": 0,
                "minimum": 70,
                "due_day": 10,
                "snowball_order": 2,
            },
            {
                "name": "Debt C",
                "balance": 900,
                "apr": 0,
                "minimum": 90,
                "due_day": 10,
                "snowball_order": 3,
            },
        ],
    }
    if scenarios_marker != "missing":
        data["scenarios"] = scenarios_marker
    if target_marker != "missing":
        data["debt_free_target"] = target_marker

    return data


def write_config(tmp_path, data=None, raw_json=None):
    config_path = tmp_path / "config.json"
    if raw_json is None:
        config_path.write_text(json.dumps(data or config_data()), encoding="utf-8")
    else:
        config_path.write_text(raw_json, encoding="utf-8")

    return config_path


def load_config(tmp_path, data=None, raw_json=None):
    return Config().load(write_config(tmp_path, data=data, raw_json=raw_json))


def scenario_config(scenario):
    return config_data(scenarios_marker=[scenario])


def test_valid_extra_payment_scenario_parses_correctly(tmp_path):
    config = load_config(
        tmp_path,
        scenario_config({"name": "Extra $75", "extra_per_paycheck": "75.00"}),
    )

    assert config.scenarios[0].name == "Extra $75"
    assert config.scenarios[0].extra_per_paycheck == Decimal("75.00")


def test_valid_savings_override_parses_correctly(tmp_path):
    config = load_config(
        tmp_path,
        scenario_config({"name": "Save Less", "savings_percentage": "0.25"}),
    )

    assert config.scenarios[0].savings_percentage_override == Decimal("0.25")


def test_valid_snowball_order_override_parses_correctly(tmp_path):
    config = load_config(
        tmp_path,
        scenario_config({"name": "Custom", "snowball_order": ["Debt C", "Debt A"]}),
    )

    assert config.scenarios[0].snowball_order_override == ["Debt C", "Debt A"]


def test_all_supported_fields_parse_together(tmp_path):
    config = load_config(
        tmp_path,
        scenario_config(
            {
                "name": "Aggressive",
                "extra_per_paycheck": "125.50",
                "savings_percentage": "0.00",
                "snowball_order": ["Debt B", "Debt A"],
            }
        ),
    )
    scenario = config.scenarios[0]

    assert scenario.extra_per_paycheck == Decimal("125.50")
    assert scenario.savings_percentage_override == Decimal("0.00")
    assert scenario.snowball_order_override == ["Debt B", "Debt A"]


@pytest.mark.parametrize("scenarios_marker", ["missing", None, []])
def test_optional_scenarios_produce_empty_alternative_list(
    tmp_path,
    scenarios_marker,
):
    config = load_config(tmp_path, config_data(scenarios_marker=scenarios_marker))

    assert config.scenarios == []


def test_old_config_without_scenarios_remains_valid(tmp_path):
    config = load_config(tmp_path, config_data())

    assert config.settings.first_paycheck == date(2026, 1, 2)
    assert [debt.name for debt in config.debts] == ["Debt A", "Debt B", "Debt C"]
    assert config.scenarios == []


def test_scenario_order_is_preserved(tmp_path):
    config = load_config(
        tmp_path,
        config_data(
            scenarios_marker=[
                {"name": "First", "extra_per_paycheck": "10.00"},
                {"name": "Second", "extra_per_paycheck": "20.00"},
            ]
        ),
    )

    assert [scenario.name for scenario in config.scenarios] == ["First", "Second"]


def test_decimal_values_retain_precision_from_numeric_json(tmp_path):
    raw_json = """
    {
      "budget": {
        "paycheck": 500,
        "first_paycheck": "2026-01-02",
        "rent_per_paycheck": 0,
        "insurance_per_paycheck": 0,
        "personal_per_paycheck": 0,
        "starting_savings": 0,
        "savings_goal": 0,
        "snowball_split": 0.5
      },
      "bills": [],
      "debts": [
        {
          "name": "Debt A",
          "balance": 500,
          "apr": 0,
          "minimum": 50,
          "due_day": 10,
          "snowball_order": 1
        }
      ],
      "scenarios": [
        {
          "name": "Precise",
          "extra_per_paycheck": 0.10,
          "savings_percentage": 0.25
        }
      ]
    }
    """

    config = load_config(tmp_path, raw_json=raw_json)

    assert config.scenarios[0].extra_per_paycheck == Decimal("0.10")
    assert config.scenarios[0].savings_percentage_override == Decimal("0.25")


def test_omitted_fields_receive_defaults(tmp_path):
    config = load_config(tmp_path, scenario_config({"name": "Defaults"}))
    scenario = config.scenarios[0]

    assert scenario.extra_per_paycheck == Decimal("0.00")
    assert scenario.savings_percentage_override is None
    assert scenario.snowball_order_override is None


@pytest.mark.parametrize(
    ("scenario", "message"),
    [
        ({}, "missing name"),
        ({"name": " "}, "blank"),
        ({"name": "Bad", "extra_per_paycheck": "-1.00"}, "negative"),
        ({"name": "Bad", "extra_per_paycheck": "nope"}, "valid decimal"),
        ({"name": "Bad", "extra_per_paycheck": "NaN"}, "finite"),
        ({"name": "Bad", "extra_per_paycheck": "Infinity"}, "finite"),
        ({"name": "Bad", "savings_percentage": "-0.01"}, "between 0 and 1"),
        ({"name": "Bad", "savings_percentage": "1.01"}, "between 0 and 1"),
        ({"name": "Bad", "snowball_order": "Debt A"}, "must be a list"),
        ({"name": "Bad", "snowball_order": [" "]}, "blank names"),
        (
            {"name": "Bad", "snowball_order": ["Debt A", "Debt A"]},
            "duplicates",
        ),
        ({"name": "Bad", "snowball_order": ["Missing"]}, "unknown debt"),
    ],
)
def test_invalid_scenario_configuration_is_rejected(tmp_path, scenario, message):
    with pytest.raises(ValueError, match=message):
        load_config(tmp_path, scenario_config(scenario))


def test_duplicate_scenario_names_are_rejected_case_insensitively(tmp_path):
    data = config_data(
        scenarios_marker=[
            {"name": "Extra", "extra_per_paycheck": "10.00"},
            {"name": "extra", "extra_per_paycheck": "20.00"},
        ]
    )

    with pytest.raises(ValueError, match="unique"):
        load_config(tmp_path, data)


@pytest.mark.parametrize(
    ("scenarios_marker", "message"),
    [
        ("not-list", "scenarios must be a list"),
        ([1], "must be an object"),
    ],
)
def test_invalid_scenarios_container_is_rejected(
    tmp_path,
    scenarios_marker,
    message,
):
    with pytest.raises(ValueError, match=message):
        load_config(tmp_path, config_data(scenarios_marker=scenarios_marker))


def test_run_orchestration_uses_configured_scenarios(tmp_path):
    config = load_config(
        tmp_path,
        config_data(
            scenarios_marker=[
                {"name": "Configured $20", "extra_per_paycheck": "20.00"},
                {"name": "Configured $40", "extra_per_paycheck": "40.00"},
            ]
        ),
    )

    comparison = build_scenario_comparison(config)

    assert comparison.baseline.name == "Baseline"
    assert [scenario.name for scenario in comparison.scenarios] == [
        "Configured $20",
        "Configured $40",
    ]


def test_baseline_only_workbook_generation_succeeds(tmp_path):
    config = load_config(tmp_path, config_data(scenarios_marker=[]))
    comparison = build_scenario_comparison(config)
    workbook_path = tmp_path / "baseline_only.xlsx"

    ExcelWriter(workbook_path).write([], None, comparison)

    workbook = load_workbook(workbook_path)
    sheet = workbook["Scenario Comparison"]
    assert sheet["A4"].value == "Baseline"
    assert sheet["A5"].value is None
    assert workbook["Dashboard"]["B4"].value == "Baseline"
    workbook.close()


def test_config_load_mapping_accepts_saved_snapshot_withdrawal_date():
    """Saved config snapshots can be loaded without rewriting withdrawal dates."""
    data = config_data()
    data["savings_plan"] = {
        "deadline_priority_enabled": False,
        "goals": [],
        "withdrawals": [
            {
                "name": "Saved withdrawal",
                "withdrawal_date": "2026-01-16",
                "amount": "25.00",
            }
        ],
    }

    config = Config().load_mapping(data)

    assert config.savings_plan.withdrawals[0].withdrawal_date == date(2026, 1, 16)


def test_config_load_mapping_accepts_legacy_withdrawal_date_key():
    """Older config-shaped snapshots still load the legacy withdrawal date key."""
    data = config_data()
    data["savings_plan"] = {
        "deadline_priority_enabled": False,
        "goals": [],
        "withdrawals": [
            {
                "name": "Legacy withdrawal",
                "date": "2026-01-16",
                "amount": "25.00",
            }
        ],
    }

    config = Config().load_mapping(data)

    assert config.savings_plan.withdrawals[0].withdrawal_date == date(2026, 1, 16)


def test_excel_contains_only_baseline_plus_configured_scenarios_in_order(tmp_path):
    config = load_config(
        tmp_path,
        config_data(
            scenarios_marker=[
                {"name": "Configured $20", "extra_per_paycheck": "20.00"},
                {"name": "Configured $40", "extra_per_paycheck": "40.00"},
            ]
        ),
    )
    comparison = build_scenario_comparison(config)
    workbook_path = tmp_path / "configured_scenarios.xlsx"

    ExcelWriter(workbook_path).write([], None, comparison)

    workbook = load_workbook(workbook_path)
    sheet = workbook["Scenario Comparison"]
    assert [sheet.cell(row=row, column=1).value for row in range(4, 7)] == [
        "Baseline",
        "Configured $20",
        "Configured $40",
    ]
    assert sheet["A7"].value is None
    workbook.close()


def test_missing_target_section_is_disabled(tmp_path):
    config = load_config(tmp_path, config_data())

    assert config.debt_free_target.enabled is False


def test_disabled_target_section_is_disabled(tmp_path):
    config = load_config(
        tmp_path,
        config_data(target_marker={"enabled": False}),
    )

    assert config.debt_free_target.enabled is False


def test_enabled_valid_target_section_parses(tmp_path):
    config = load_config(
        tmp_path,
        config_data(
            target_marker={
                "enabled": True,
                "target_date": "2026-01-30",
                "maximum_extra_per_paycheck": "250.00",
                "precision": "0.10",
                "maximum_iterations": 25,
            }
        ),
    )

    assert config.debt_free_target.enabled is True
    assert config.debt_free_target.target_date == date(2026, 1, 30)
    assert config.debt_free_target.maximum_extra_per_paycheck == Decimal("250.00")
    assert config.debt_free_target.precision == Decimal("0.10")
    assert config.debt_free_target.maximum_iterations == 25


@pytest.mark.parametrize(
    ("target_marker", "message"),
    [
        ({"enabled": True}, "target_date"),
        ({"enabled": True, "target_date": "01/30/2026"}, "YYYY-MM-DD"),
        ({"enabled": "yes", "target_date": "2026-01-30"}, "boolean"),
        (
            {
                "enabled": True,
                "target_date": "2026-01-30",
                "maximum_extra_per_paycheck": "-0.01",
            },
            "negative",
        ),
        (
            {
                "enabled": True,
                "target_date": "2026-01-30",
                "maximum_extra_per_paycheck": "NaN",
            },
            "finite",
        ),
        (
            {
                "enabled": True,
                "target_date": "2026-01-30",
                "precision": "0.00",
            },
            "greater than zero",
        ),
        (
            {
                "enabled": True,
                "target_date": "2026-01-30",
                "maximum_iterations": 0,
            },
            "positive",
        ),
        (
            {
                "enabled": True,
                "target_date": "2026-01-30",
                "maximum_iterations": "ten",
            },
            "integer",
        ),
    ],
)
def test_invalid_target_configuration_is_rejected(tmp_path, target_marker, message):
    with pytest.raises(ValueError, match=message):
        load_config(tmp_path, config_data(target_marker=target_marker))


def test_backward_compatibility_with_old_config_has_disabled_target(tmp_path):
    config = load_config(tmp_path, config_data())

    assert config.scenarios == []
    assert config.debt_free_target.enabled is False


def test_run_orchestration_returns_none_when_target_disabled(tmp_path):
    config = load_config(tmp_path, config_data(target_marker={"enabled": False}))

    assert build_debt_free_target_result(config) is None


def test_run_orchestration_calculates_when_target_enabled(tmp_path):
    config = load_config(
        tmp_path,
        config_data(
            target_marker={
                "enabled": True,
                "target_date": "2026-01-16",
                "maximum_extra_per_paycheck": "1000.00",
                "precision": "0.01",
            }
        ),
    )

    result = build_debt_free_target_result(config)

    assert result.target_met is True
    assert result.calculation_status in {
        DebtFreeTargetStatus.TARGET_MET,
        DebtFreeTargetStatus.ALREADY_ON_TRACK,
    }
