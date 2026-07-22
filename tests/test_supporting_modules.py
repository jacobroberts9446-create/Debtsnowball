import json
import sqlite3
from contextlib import closing
from datetime import date
from decimal import Decimal
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from app.budget_engine import DebtBalance, PayPeriodSummary
from app.calendar_engine import CalendarEngine
from app.config import Config
from app.database import Database
from app.excel_writer import ExcelWriter
from app.models import (
    Bill,
    BudgetSettings,
    Debt,
    DebtPayoffForecast,
    ForecastPeriod,
    ForecastSummary,
    Savings,
)
from app.money import money


def test_calendar_engine_generates_biweekly_periods():
    settings = SimpleNamespace(first_paycheck=date(2026, 1, 2))

    periods = CalendarEngine(settings).generate(date(2026, 1, 30))

    assert [period.pay_date for period in periods] == [
        date(2026, 1, 2),
        date(2026, 1, 16),
        date(2026, 1, 30),
    ]
    assert periods[0].start_date == date(2026, 1, 2)
    assert periods[0].end_date == date(2026, 1, 15)


def test_config_loads_budget_bills_and_sorted_debts(tmp_path):
    config_path = tmp_path / "config.json"
    config_path.write_text(
        json.dumps(
            {
                "budget": {
                    "paycheck": 1000,
                    "first_paycheck": "2026-01-02",
                    "rent_per_paycheck": 100,
                    "insurance_per_paycheck": 50,
                    "personal_per_paycheck": 200,
                    "starting_savings": 300,
                    "savings_goal": 1000,
                    "snowball_split": 0.5,
                },
                "bills": [{"name": "Phone", "amount": 100, "due_day": 10}],
                "debts": [
                    {
                        "name": "Second",
                        "balance": 200,
                        "apr": 0,
                        "minimum": 20,
                        "due_day": 12,
                        "snowball_order": 2,
                    },
                    {
                        "name": "First",
                        "balance": 100,
                        "apr": 0,
                        "minimum": 10,
                        "due_day": 11,
                        "snowball_order": 1,
                    },
                ],
            }
        )
    )

    config = Config()
    config.load(config_path)

    assert config.settings.first_paycheck == date(2026, 1, 2)
    assert config.bills[0].name == "Phone"
    assert [debt.name for debt in config.debts] == ["First", "Second"]


def test_config_raises_for_missing_file(tmp_path):
    with pytest.raises(FileNotFoundError):
        Config().load(tmp_path / "missing.json")


def test_database_saves_paychecks_and_debts(tmp_path):
    db_path = tmp_path / "plan.sqlite"
    paychecks = [
        SimpleNamespace(
            pay_date=date(2026, 1, 2),
            income=1000,
            bills_paid=100,
            debt_minimums=50,
            snowball_payment=400,
            savings_added=450,
            checking_remaining=0,
            notes=["ok"],
        )
    ]
    debts = [
        {
            "name": "Card",
            "balance": 500,
            "apr": 20,
            "minimum": 50,
            "total_paid": 100,
            "total_interest_paid": 10,
            "status": "active",
        }
    ]

    Database(db_path).save_plan(paychecks, debts)

    with closing(sqlite3.connect(db_path)) as conn:
        assert conn.execute("SELECT COUNT(*) FROM paychecks").fetchone()[0] == 1
        assert conn.execute("SELECT COUNT(*) FROM debts").fetchone()[0] == 1


def test_database_real_boundary_reloads_money_as_normalized_decimal(tmp_path):
    db_path = tmp_path / "plan.sqlite"
    paychecks = [
        SimpleNamespace(
            pay_date=date(2026, 1, 2),
            income=Decimal("1500.00"),
            bills_paid=Decimal("215.00"),
            debt_minimums=Decimal("0.01"),
            snowball_payment=Decimal("568.50"),
            savings_added=Decimal("2400.00"),
            checking_remaining=Decimal("-0.00"),
            notes=[],
        )
    ]
    debts = [
        {
            "name": "Card",
            "balance": Decimal("999999999.99"),
            "apr": Decimal("26"),
            "minimum": Decimal("0.02"),
            "total_paid": Decimal("215.00"),
            "total_interest_paid": Decimal("23.46"),
            "status": "active",
        }
    ]
    database = Database(db_path)

    database.save_plan(paychecks, debts)
    saved_paycheck = database.load_paychecks()[0]
    saved_debt = database.load_debts()[0]

    assert saved_paycheck["income"] == Decimal("1500.00")
    assert saved_paycheck["bills_paid"] == Decimal("215.00")
    assert saved_paycheck["debt_minimums"] == Decimal("0.01")
    assert saved_paycheck["snowball_payment"] == Decimal("568.50")
    assert saved_paycheck["savings_added"] == Decimal("2400.00")
    assert saved_paycheck["checking_remaining"] == Decimal("0.00")
    assert not saved_paycheck["checking_remaining"].is_signed()
    assert saved_debt["balance"] == Decimal("999999999.99")
    assert saved_debt["minimum"] == Decimal("0.02")
    assert saved_debt["total_interest_paid"] == Decimal("23.46")


def test_excel_writer_creates_dashboard_tables_and_charts(tmp_path):
    workbook_path = tmp_path / "plan.xlsx"
    summaries = [
        PayPeriodSummary(
            pay_date=date(2026, 1, 2),
            start_date=date(2026, 1, 2),
            end_date=date(2026, 1, 15),
            income=1000,
            bills_paid=100,
            debt_minimums=50,
            snowball_payment=400,
            savings_contribution=450,
            savings_balance=950,
            savings_goal=1000,
            remaining_cash=0,
            active_debt_balances=[
                DebtBalance("Card", 500, 50, "active"),
                DebtBalance("Loan", 1000, 100, "active"),
            ],
            paid_off_debts=[],
        ),
        PayPeriodSummary(
            pay_date=date(2026, 1, 16),
            start_date=date(2026, 1, 16),
            end_date=date(2026, 1, 29),
            income=1000,
            bills_paid=100,
            debt_minimums=50,
            snowball_payment=850,
            savings_contribution=0,
            savings_balance=1000,
            savings_goal=1000,
            remaining_cash=0,
            active_debt_balances=[DebtBalance("Loan", 650, 100, "active")],
            paid_off_debts=[DebtBalance("Card", 0, 50, "paid")],
        ),
    ]
    forecast = ForecastSummary(
        forecast_start_date=date(2026, 1, 2),
        forecast_end_date=date(2026, 1, 16),
        debt_free_date=date(2026, 1, 16),
        savings_goal_date=date(2026, 1, 16),
        starting_debt=Decimal("1500.00"),
        total_interest_paid=Decimal("25.00"),
        total_minimum_payments=Decimal("100.00"),
        total_snowball_payments=Decimal("1250.00"),
        ending_savings=Decimal("1000.00"),
        remaining_debt=Decimal("650.00"),
        completed=True,
        debt_payoffs=[
            DebtPayoffForecast(
                debt_name="Card",
                starting_balance=Decimal("500.00"),
                payoff_date=date(2026, 1, 16),
                total_interest_paid=Decimal("25.00"),
                total_paid=Decimal("525.00"),
            )
        ],
        periods=[
            ForecastPeriod(
                paycheck_date=date(2026, 1, 2),
                total_debt_balance=Decimal("1500.00"),
                savings_balance=Decimal("950.00"),
                interest_paid=Decimal("10.00"),
                minimums_paid=Decimal("50.00"),
                snowball_paid=Decimal("400.00"),
            ),
            ForecastPeriod(
                paycheck_date=date(2026, 1, 16),
                total_debt_balance=Decimal("650.00"),
                savings_balance=Decimal("1000.00"),
                interest_paid=Decimal("15.00"),
                minimums_paid=Decimal("50.00"),
                snowball_paid=Decimal("850.00"),
            ),
        ],
    )

    ExcelWriter(workbook_path).write(summaries, forecast)

    workbook = load_workbook(workbook_path)
    assert workbook.sheetnames == [
        "Dashboard",
        "Pay Period Summaries",
        "Active Debts",
        "Paid-Off Debts",
        "Savings Progress",
        "Forecast",
        "Scenario Comparison",
    ]
    assert len(workbook["Dashboard"]._charts) == 2
    assert workbook["Dashboard"]["A12"].value == "Estimated Debt-Free Date"
    assert money(workbook["Dashboard"]["B14"].value) == Decimal("25.00")
    assert workbook["Dashboard"].freeze_panes == "A4"
    assert workbook["Dashboard"]._charts[0].x_axis.number_format.formatCode == "@"
    assert workbook["Dashboard"]._charts[0].x_axis.tickLblSkip == 1
    assert workbook["Pay Period Summaries"].auto_filter.ref == "A1:W3"
    assert workbook["Pay Period Summaries"]["W2"].value == "Jan 02"
    assert workbook["Pay Period Summaries"].column_dimensions["W"].hidden is True
    assert workbook["Active Debts"].max_row == 4
    assert workbook["Paid-Off Debts"].max_row == 2
    assert workbook["Savings Progress"]["C3"].value == 1000
    assert workbook["Forecast"]["A1"].value == "Forecast"
    assert workbook["Forecast"]["A4"].value == "Estimated Debt-Free Date"
    assert workbook["Forecast"]["A14"].value == "Debt"
    assert workbook["Forecast"]["A19"].value == "Paycheck Date"
    assert workbook["Forecast"].auto_filter.ref == "A19:E21"
    assert len(workbook["Forecast"]._charts) == 2
    assert workbook["Forecast"]._charts[0].x_axis.number_format.formatCode == "@"
    assert workbook["Forecast"]._charts[0].x_axis.tickLblSkip == 1
    assert workbook["Forecast"]["F20"].value == "Jan 02"
    assert workbook["Forecast"].column_dimensions["F"].hidden is True
    workbook.close()


def test_excel_writer_handles_empty_forecast(tmp_path):
    workbook_path = tmp_path / "empty_forecast.xlsx"

    ExcelWriter(workbook_path).write([], None)

    workbook = load_workbook(workbook_path)
    assert "Forecast" in workbook.sheetnames
    assert "Scenario Comparison" in workbook.sheetnames
    assert workbook["Forecast"]["A1"].value == "Forecast"
    assert workbook["Forecast"]["A4"].value == "Estimated Debt-Free Date"
    assert len(workbook["Forecast"]._charts) == 0
    assert workbook["Dashboard"]["A1"].value == "DebtSnowball Dashboard"
    workbook.close()


def test_excel_writer_normalizes_negative_zero_money_cells(tmp_path):
    workbook_path = tmp_path / "negative_zero.xlsx"
    summaries = [
        PayPeriodSummary(
            pay_date=date(2026, 1, 2),
            start_date=date(2026, 1, 2),
            end_date=date(2026, 1, 15),
            income=Decimal("100.00"),
            bills_paid=Decimal("-0.00"),
            debt_minimums=Decimal("-0.00"),
            snowball_payment=Decimal("-0.00"),
            savings_contribution=Decimal("-0.00"),
            savings_balance=Decimal("-0.00"),
            savings_goal=Decimal("100.00"),
            remaining_cash=Decimal("-0.00"),
            active_debt_balances=[
                DebtBalance("Card", Decimal("-0.00"), Decimal("-0.00"), "paid")
            ],
            paid_off_debts=[
                DebtBalance("Card", Decimal("-0.00"), Decimal("-0.00"), "paid")
            ],
        )
    ]

    ExcelWriter(workbook_path).write(summaries, None)

    workbook = load_workbook(workbook_path, data_only=True)
    pay_period = workbook["Pay Period Summaries"]
    active_debts = workbook["Active Debts"]
    paid_off_debts = workbook["Paid-Off Debts"]
    assert pay_period["E2"].value == 0
    assert pay_period["I2"].value == 0
    assert pay_period["N2"].value == 0
    assert pay_period["R2"].value == 0
    assert pay_period["T2"].value == 0
    assert active_debts["C2"].value == 0
    assert active_debts["D2"].value == 0
    assert paid_off_debts["C2"].value == 0
    assert paid_off_debts["D2"].value == 0
    workbook.close()


def test_excel_writer_uses_sparse_labels_for_long_forecast_charts(tmp_path):
    workbook_path = tmp_path / "long_forecast.xlsx"
    periods = [
        ForecastPeriod(
            paycheck_date=date(2026 + index // 26, 1 + (index % 12), 1),
            total_debt_balance=Decimal("1000.00") - Decimal(index),
            savings_balance=Decimal(index),
            interest_paid=Decimal("1.00"),
            minimums_paid=Decimal("10.00"),
            snowball_paid=Decimal("20.00"),
        )
        for index in range(40)
    ]
    forecast = ForecastSummary(
        forecast_start_date=date(2026, 1, 1),
        forecast_end_date=date(2027, 4, 1),
        debt_free_date=None,
        savings_goal_date=None,
        starting_debt=Decimal("1000.00"),
        total_interest_paid=Decimal("40.00"),
        total_minimum_payments=Decimal("400.00"),
        total_snowball_payments=Decimal("800.00"),
        ending_savings=Decimal("40.00"),
        remaining_debt=Decimal("960.00"),
        completed=False,
        periods=periods,
    )

    ExcelWriter(workbook_path).write([], forecast)

    workbook = load_workbook(workbook_path)
    chart_axis = workbook["Forecast"]._charts[0].x_axis
    assert chart_axis.number_format.formatCode == "@"
    assert chart_axis.tickLblSkip == 5
    assert chart_axis.tickMarkSkip == 5
    assert workbook["Forecast"]["F20"].value == "Jan 2026"
    assert workbook["Forecast"].column_dimensions["F"].hidden is True
    workbook.close()


def test_model_validation_and_savings_goal():
    with pytest.raises(ValueError):
        Bill("Bad Bill", 10, 0)

    with pytest.raises(ValueError):
        Debt("Bad Balance", -1, 0, 10, 1, 1)

    with pytest.raises(ValueError):
        Debt("Bad Apr", 100, -1, 10, 1, 1)

    with pytest.raises(ValueError):
        Debt("Bad Minimum", 100, 0, -1, 1, 1)

    with pytest.raises(ValueError):
        Debt("Bad Due Day", 100, 0, 10, 32, 1)

    with pytest.raises(ValueError):
        BudgetSettings(100, date(2026, 1, 1), 0, 0, 0, 0, 100, 1.5)

    with pytest.raises(ValueError):
        BudgetSettings(-1, date(2026, 1, 1), 0, 0, 0, 0, 100, 0.5)

    savings = Savings(50, 100)
    assert not savings.goal_met
    savings.add(50)
    assert savings.goal_met
