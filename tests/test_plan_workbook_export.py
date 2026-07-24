"""Tests for workbook export workflows."""

from datetime import date
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace

from openpyxl import load_workbook

from app.workflows import workbook_export
import run


def output_text(output: list[str]) -> str:
    """Join captured console output for readable substring assertions."""
    return "\n".join(output)


class FakePlanHistoryService:
    """Small fake service for workbook export workflow tests."""

    def __init__(self, plans=None) -> None:
        self.plans = plans or []
        self.versions_by_plan = {}

    def list_plans(self):
        return self.plans

    def get_plan(self, plan_id):
        for plan in self.plans:
            if plan.id == plan_id:
                return plan
        raise ValueError(f"plan {plan_id} was not found.")

    def get_plan_version(self, version_id):
        for versions in self.versions_by_plan.values():
            for version in versions:
                if version.id == version_id:
                    return version
        raise ValueError(f"plan version {version_id} was not found.")

    def list_plan_versions(self, plan_id):
        return self.versions_by_plan.get(plan_id, [])


def test_config_generation_action_remains_available_for_developer_workflows() -> None:
    """Config generation remains callable outside the interactive menu."""
    calls = []

    result = workbook_export.run_generate_budget_plan_action(
        lambda: calls.append("generated")
    )

    assert result is False
    assert calls == ["generated"]


def test_config_generation_action_reports_recoverable_error() -> None:
    """Recoverable generation errors are still shown by the developer action."""
    output = []

    result = workbook_export.run_generate_budget_plan_action(
        lambda: (_ for _ in ()).throw(ValueError("bad config")),
        output_func=output.append,
    )

    assert result is False
    assert "Error: bad config" in output


def test_config_generation_action_catches_system_exit() -> None:
    """SystemExit from generation is reported by the developer action."""
    output = []

    result = workbook_export.run_generate_budget_plan_action(
        lambda: (_ for _ in ()).throw(SystemExit(2)),
        output_func=output.append,
    )

    assert result is False
    assert "Error: Plan generation stopped unexpectedly: 2" in output


def test_generate_budget_plan_prints_full_workbook_path(monkeypatch) -> None:
    """Config generation prints a clear workbook path for users."""
    output = []

    class FakeConfig:
        settings = SimpleNamespace()
        scenarios = []
        debt_free_target = SimpleNamespace(enabled=False)

        def load(self):
            return self

    class FakeCalendar:
        def __init__(self, _settings) -> None:
            pass

        def generate(self, _end_date):
            return []

    class FakeBudget:
        def __init__(self, _config) -> None:
            pass

        def build_plan(self, _periods):
            return []

    class FakeForecast:
        def __init__(self, _config) -> None:
            pass

        def forecast(self):
            return SimpleNamespace()

    class FakeWriter:
        def write(self, *_args):
            return Path(
                "C:/Users/Test/AppData/Local/DebtSnowball/output/debtsnowball_plan.xlsx"
            )

    monkeypatch.setattr(workbook_export, "Config", FakeConfig)
    monkeypatch.setattr(workbook_export, "CalendarEngine", FakeCalendar)
    monkeypatch.setattr(workbook_export, "BudgetEngine", FakeBudget)
    monkeypatch.setattr(workbook_export, "ForecastEngine", FakeForecast)
    monkeypatch.setattr(workbook_export, "ExcelWriter", lambda: FakeWriter())
    monkeypatch.setattr(
        workbook_export, "build_scenario_comparison", lambda _config: None
    )
    monkeypatch.setattr(
        workbook_export,
        "print_success",
        lambda message: output.append(f"Success: {message}"),
    )
    monkeypatch.setattr(
        workbook_export,
        "print",
        lambda value="": output.append(value),
        raising=False,
    )

    workbook_export.generate_budget_plan()

    text = output_text([str(item) for item in output])
    assert "Success: Excel workbook created:" in text
    assert (
        "C:\\Users\\Test\\AppData\\Local\\DebtSnowball\\output\\debtsnowball_plan.xlsx"
        in text
    )


def test_generate_budget_plan_prints_detailed_pay_period_summary(monkeypatch) -> None:
    """Legacy config generation prints all detailed pay-period lines."""
    output = []
    workbook_path = Path(
        "C:/Users/Test/AppData/Local/DebtSnowball/output/debtsnowball_plan.xlsx"
    )
    summary = SimpleNamespace(
        start_date=date(2026, 1, 2),
        end_date=date(2026, 1, 15),
        income=Decimal("1000.00"),
        bills_paid=Decimal("100.00"),
        debt_minimums=Decimal("25.00"),
        active_savings_goal_name="Emergency Fund",
        available_after_required_payments=Decimal("875.00"),
        snowball_reduction=Decimal("50.00"),
        personal_expense_reduction=Decimal("10.00"),
        projected_savings_shortfall=Decimal("5.00"),
        savings_contribution=Decimal("60.00"),
        snowball_payment=Decimal("40.00"),
        remaining_cash=Decimal("0.00"),
        active_debt_balances=[SimpleNamespace(name="Card", balance=Decimal("75.00"))],
        paid_off_debts=[SimpleNamespace(name="Loan")],
    )

    class FakeConfig:
        def load(self):
            return self

    class FakeWriter:
        def write(self, *_args):
            return workbook_path

    monkeypatch.setattr(workbook_export, "Config", FakeConfig)
    monkeypatch.setattr(
        workbook_export,
        "build_workbook_outputs",
        lambda _config: ([summary], "forecast", None, None),
    )
    monkeypatch.setattr(workbook_export, "ExcelWriter", lambda: FakeWriter())
    monkeypatch.setattr(
        workbook_export,
        "print_success",
        lambda message: output.append(f"Success: {message}"),
    )
    monkeypatch.setattr(
        workbook_export,
        "print_application_banner",
        lambda version: output.append(f"Banner {version}"),
    )
    monkeypatch.setattr(
        workbook_export,
        "print_section_header",
        lambda title: output.append(title),
    )
    monkeypatch.setattr(
        workbook_export,
        "print",
        lambda value="": output.append(value),
        raising=False,
    )

    workbook_export.generate_budget_plan()

    text = output_text([str(item) for item in output])
    assert "Pay Period: Jan 02, 2026 to Jan 15, 2026" in text
    assert "Income:          $1,000.00" in text
    assert "Savings Goal:    Emergency Fund" in text
    assert "Personal Expense Reduction: $10.00" in text
    assert "Projected Savings Shortfall: $5.00" in text
    assert "Card            $75.00" in text
    assert "Paid Off: Loan" in text
    assert str(workbook_path) in text


def test_saved_plan_workbook_generation_displays_path(monkeypatch, tmp_path) -> None:
    """Plan details can generate a workbook and display the full path."""
    choices = iter([""])
    output = []
    workbook_path = tmp_path / "output" / "debtsnowball_plan.xlsx"
    service = FakePlanHistoryService(
        [
            SimpleNamespace(
                id=4,
                name="Current Plan",
                description="",
                updated_at="",
                current_version_id=21,
            )
        ]
    )
    service.versions_by_plan[4] = [SimpleNamespace(id=21, version_number=1)]
    monkeypatch.setattr(
        workbook_export, "config_from_plan_version", lambda _version: object()
    )
    monkeypatch.setattr(
        workbook_export,
        "build_workbook_outputs",
        lambda _config: (["summary"], "forecast", "scenarios", "target"),
    )

    class FakeWriter:
        def write(self, *_args):
            workbook_path.parent.mkdir(parents=True, exist_ok=True)
            workbook_path.write_bytes(b"workbook")
            return workbook_path

    monkeypatch.setattr(workbook_export, "ExcelWriter", lambda: FakeWriter())

    workbook_export.generate_saved_plan_workbook_action(
        service,
        service.plans[0],
        input_func=lambda _prompt: next(choices),
        output_func=output.append,
    )

    assert "Success: Workbook created successfully." in output
    assert "Location:" in output
    assert str(workbook_path) in output


def test_saved_plan_workbook_generation_failure_does_not_claim_success(
    monkeypatch,
    tmp_path,
) -> None:
    """A missing workbook file reports failure without claiming success."""
    choices = iter([""])
    output = []
    missing_path = tmp_path / "output" / "missing.xlsx"
    service = FakePlanHistoryService(
        [
            SimpleNamespace(
                id=4,
                name="Current Plan",
                description="",
                updated_at="",
                current_version_id=21,
            )
        ]
    )
    service.versions_by_plan[4] = [SimpleNamespace(id=21, version_number=1)]
    monkeypatch.setattr(
        workbook_export, "config_from_plan_version", lambda _version: object()
    )
    monkeypatch.setattr(
        workbook_export,
        "build_workbook_outputs",
        lambda _config: (["summary"], "forecast", "scenarios", "target"),
    )

    class MissingWriter:
        def write(self, *_args):
            return missing_path

    monkeypatch.setattr(workbook_export, "ExcelWriter", lambda: MissingWriter())

    workbook_export.generate_saved_plan_workbook_action(
        service,
        service.plans[0],
        input_func=lambda _prompt: next(choices),
        output_func=output.append,
    )

    text = output_text(output)
    assert "Error: workbook was not created:" in text
    assert "Workbook created successfully." not in text


def test_saved_plan_workbook_generation_handles_no_versions() -> None:
    """A saved plan with no versions reports a handled error and returns."""
    choices = iter([""])
    output = []
    service = FakePlanHistoryService([SimpleNamespace(id=4, name="No Versions")])

    result = workbook_export.generate_saved_plan_workbook_action(
        service,
        service.plans[0],
        input_func=lambda _prompt: next(choices),
        output_func=output.append,
    )

    text = output_text(output)
    assert result is False
    assert "Error: selected plan does not have any saved versions." in text
    assert "Workbook created successfully." not in text


def test_saved_plan_workbook_generation_handles_writer_exception(monkeypatch) -> None:
    """Writer exceptions are reported without claiming success."""
    choices = iter([""])
    output = []
    service = FakePlanHistoryService(
        [SimpleNamespace(id=4, name="Plan", current_version_id=21)]
    )
    service.versions_by_plan[4] = [SimpleNamespace(id=21, version_number=1)]
    monkeypatch.setattr(
        workbook_export, "config_from_plan_version", lambda _version: object()
    )
    monkeypatch.setattr(
        workbook_export,
        "build_workbook_outputs",
        lambda _config: (["summary"], "forecast", "scenarios", "target"),
    )

    class FailingWriter:
        def write(self, *_args):
            raise OSError("output directory unavailable")

    monkeypatch.setattr(workbook_export, "ExcelWriter", lambda: FailingWriter())

    result = workbook_export.generate_saved_plan_workbook_action(
        service,
        service.plans[0],
        input_func=lambda _prompt: next(choices),
        output_func=output.append,
    )

    text = output_text(output)
    assert result is False
    assert "Error: output directory unavailable" in text
    assert "Workbook created successfully." not in text


def test_verify_workbook_created_rejects_directory_and_empty_file(tmp_path) -> None:
    """Workbook verification requires a non-empty regular file."""
    directory = tmp_path / "folder.xlsx"
    directory.mkdir()
    empty_file = tmp_path / "empty.xlsx"
    empty_file.write_bytes(b"")
    workbook_file = tmp_path / "workbook.xlsx"
    workbook_file.write_bytes(b"not empty")

    try:
        workbook_export.verify_workbook_created(directory)
    except RuntimeError as exc:
        assert "workbook path is not a file:" in str(exc)
    else:
        raise AssertionError("directory path should fail workbook verification")

    try:
        workbook_export.verify_workbook_created(empty_file)
    except RuntimeError as exc:
        assert "workbook file is empty:" in str(exc)
    else:
        raise AssertionError("empty workbook should fail verification")

    assert workbook_export.verify_workbook_created(workbook_file) == workbook_file


def test_latest_plan_version_uses_last_version_when_current_is_missing() -> None:
    """Latest version falls back to the final listed version for older plan rows."""
    service = FakePlanHistoryService([SimpleNamespace(id=4, name="Plan")])
    service.versions_by_plan[4] = [
        SimpleNamespace(id=20, version_number=1),
        SimpleNamespace(id=21, version_number=2),
    ]

    version = workbook_export.latest_plan_version(service, service.plans[0])

    assert version.id == 21


def test_latest_plan_version_returns_current_version_when_present() -> None:
    """Latest version prefers the explicit current version ID when available."""
    service = FakePlanHistoryService(
        [SimpleNamespace(id=4, name="Plan", current_version_id=20)]
    )
    service.versions_by_plan[4] = [
        SimpleNamespace(id=20, version_number=1),
        SimpleNamespace(id=21, version_number=2),
    ]

    version = workbook_export.latest_plan_version(service, service.plans[0])

    assert version.id == 20


def test_config_from_plan_version_rejects_invalid_json() -> None:
    """Invalid saved snapshots are surfaced as clear validation errors."""
    version = SimpleNamespace(id=21, config_snapshot="{")

    try:
        workbook_export.config_from_plan_version(version)
    except ValueError as exc:
        assert "saved plan version 21 has an invalid config snapshot:" in str(exc)
    else:
        raise AssertionError("invalid JSON should fail config snapshot loading")


def test_config_from_plan_version_rejects_missing_snapshot() -> None:
    """Missing saved config snapshots are surfaced as validation errors."""
    version = SimpleNamespace(id=21)

    try:
        workbook_export.config_from_plan_version(version)
    except ValueError as exc:
        assert "saved plan version 21 has an invalid config snapshot:" in str(exc)
    else:
        raise AssertionError("missing config_snapshot should fail config loading")


def test_debt_free_target_helper_runs_when_enabled(monkeypatch) -> None:
    """The optional target calculator is invoked only when enabled."""
    calls = []
    config = SimpleNamespace(debt_free_target=SimpleNamespace(enabled=True))

    class FakeCalculator:
        def __init__(self, received_config, received_request) -> None:
            calls.append((received_config, received_request))

        def calculate(self):
            return "target-result"

    monkeypatch.setattr(workbook_export, "DebtFreeTargetCalculator", FakeCalculator)

    result = workbook_export.build_debt_free_target_result(config)

    assert result == "target-result"
    assert calls == [(config, config.debt_free_target)]


def test_history_report_export_writes_expected_sections(monkeypatch, tmp_path) -> None:
    """History-report export delegates all persisted sections to ExcelWriter."""
    output = []
    report_path = tmp_path / "history.xlsx"
    written = {}
    service = SimpleNamespace(
        get_plan=lambda plan_id: SimpleNamespace(id=plan_id, name="Plan"),
        history_report_rows=lambda _plan_id: {
            "forecast_snapshots": ["snapshot"],
            "debt_history": ["debt"],
            "savings_history": ["savings"],
            "warnings": ["warning"],
        },
        list_plan_versions=lambda _plan_id: ["version"],
        compare_forecast_to_actual=lambda _plan_id: "actual",
        compare_forecast_to_actual_periods=lambda _plan_id: ["period"],
    )

    class FakeWriter:
        def __init__(self, path) -> None:
            written["path"] = path

        def write_history_report(self, plan, versions, **kwargs):
            written["plan"] = plan
            written["versions"] = versions
            written["kwargs"] = kwargs
            return report_path

    monkeypatch.setattr(workbook_export, "ExcelWriter", FakeWriter)
    monkeypatch.setattr(
        workbook_export,
        "print",
        lambda value="": output.append(value),
        raising=False,
    )

    workbook_export.write_history_report(service, 4, str(report_path))

    assert written["path"] == str(report_path)
    assert written["plan"].id == 4
    assert written["versions"] == ["version"]
    assert written["kwargs"] == {
        "actual_comparison": "actual",
        "forecast_snapshots": ["snapshot"],
        "actual_periods": ["period"],
        "debt_history": ["debt"],
        "savings_history": ["savings"],
        "warnings": ["warning"],
    }
    assert output == [f"History report created: {report_path}"]


def test_incomplete_saved_config_reports_clear_workbook_failure() -> None:
    """Malformed saved config snapshots report the real failure without success."""
    choices = iter([""])
    output = []
    service = FakePlanHistoryService(
        [
            SimpleNamespace(
                id=4,
                name="Current Plan",
                description="",
                updated_at="",
                current_version_id=21,
            )
        ]
    )
    service.versions_by_plan[4] = [
        SimpleNamespace(id=21, version_number=1, config_snapshot="{}")
    ]

    workbook_export.generate_saved_plan_workbook_action(
        service,
        service.plans[0],
        input_func=lambda _prompt: next(choices),
        output_func=output.append,
    )

    text = output_text(output)
    assert "Error:" in text
    assert "Workbook created successfully." not in text


def test_guided_saved_plan_generates_physical_workbook(monkeypatch, tmp_path) -> None:
    """A user-created saved plan can generate a real workbook from Saved Plans."""
    monkeypatch.setenv("DEBTSNOWBALL_DATA_DIR", str(tmp_path))
    workbook_path = tmp_path / "output" / "debtsnowball_plan.xlsx"
    assert not workbook_path.parent.exists()
    assert not workbook_path.exists()

    first_run_choices = iter(
        [
            "1",
            "Real User Plan",
            "2",
            "01/02/2026",
            "1000",
            "1",
            "Debt A",
            "100",
            "0",
            "10",
            "10",
            "n",
            "4",
            "1",
            "n",
            "4",
            "1",
            "0",
            "1",
            "0",
            "0",
            "2",
            "1",
            "4",
            "",
            "",
            "5",
            "2",
            "1",
            "2",
            "",
            "8",
            "2",
            "4",
        ],
    )
    output = []

    run.run_main_menu(
        input_func=lambda _prompt: next(first_run_choices),
        output_func=output.append,
    )

    text = output_text(output)
    assert "Success: Workbook created successfully." in text
    assert str(workbook_path) in text
    assert workbook_path.exists()
    assert workbook_path.is_file()
    assert workbook_path.stat().st_size > 0
    workbook = load_workbook(workbook_path, read_only=True)
    workbook.close()
    assert "Success: Goodbye." in output

    workbook_path.write_bytes(b"stale")
    stale_mtime = workbook_path.stat().st_mtime_ns
    assert workbook_path.stat().st_size == len(b"stale")

    second_run_choices = iter(["2", "1", "2", "", "8", "2", "4"])
    second_output = []
    run.run_main_menu(
        input_func=lambda _prompt: next(second_run_choices),
        output_func=second_output.append,
    )

    second_text = output_text(second_output)
    assert "Success: Workbook created successfully." in second_text
    assert str(workbook_path) in second_text
    assert workbook_path.read_bytes() != b"stale"
    assert workbook_path.stat().st_size > 0
    assert workbook_path.stat().st_mtime_ns >= stale_mtime
    workbook = load_workbook(workbook_path, read_only=True)
    workbook.close()
