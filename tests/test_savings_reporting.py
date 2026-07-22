from copy import deepcopy
from datetime import date
from decimal import Decimal
from types import SimpleNamespace

from openpyxl import load_workbook

from app.budget_engine import BudgetEngine
from app.calendar_engine import CalendarEngine
from app.config import Config
from app.excel_writer import ExcelWriter
from app.forecast_engine import ForecastEngine
from app.models import (
    BudgetSettings,
    Debt,
    PlannedSavingsWithdrawal,
    SavingsFundingMode,
    SavingsGoalStage,
    SavingsPlan,
)
from app.money import money


def make_config():
    return SimpleNamespace(
        settings=BudgetSettings(
            paycheck=1000,
            first_paycheck=date(2026, 1, 1),
            rent_per_paycheck=0,
            insurance_per_paycheck=0,
            personal_per_paycheck=0,
            starting_savings=0,
            savings_goal=1000,
            snowball_split=0.50,
        ),
        bills=[],
        debts=[
            Debt(
                "Long debt",
                balance=100000,
                apr=0,
                minimum=0,
                due_day=10,
                snowball_order=1,
            )
        ],
        scenarios=[],
        savings_plan=SavingsPlan(
            deadline_priority_enabled=True,
            goals=[
                SavingsGoalStage(
                    name="August savings",
                    target_amount=Decimal("3900.00"),
                    start_date=date(2026, 1, 1),
                    target_date=date(2026, 8, 11),
                    funding_mode=SavingsFundingMode.DEADLINE_PRIORITY,
                ),
                SavingsGoalStage(
                    name="Replacement savings",
                    target_amount=Decimal("3000.00"),
                    start_date=date(2026, 8, 12),
                    funding_mode=SavingsFundingMode.PRIORITY_UNTIL_FUNDED,
                ),
            ],
            withdrawals=[
                PlannedSavingsWithdrawal(
                    name="August planned expense",
                    withdrawal_date=date(2026, 8, 11),
                    drain_balance=True,
                )
            ],
        ),
)


def value_for_header(sheet, header: str, row: int, header_row: int = 1):
    for cell in sheet[header_row]:
        if cell.value == header:
            return sheet.cell(row=row, column=cell.column).value
    raise AssertionError(f"Missing header: {header}")


def test_savings_plan_sections_and_dashboard_metrics(tmp_path):
    config = make_config()
    periods = CalendarEngine(config.settings).generate(date(2026, 8, 28))
    summaries = BudgetEngine(config).build_plan(periods)
    forecast = ForecastEngine(config).forecast()
    workbook_path = tmp_path / "savings_plan.xlsx"

    ExcelWriter(workbook_path).write(summaries, forecast)

    workbook = load_workbook(workbook_path)
    sheet = workbook["Savings Progress"]
    dashboard = workbook["Dashboard"]

    assert sheet["A1"].value == "Savings Plan Summary"
    assert sheet["A4"].value == "August savings"
    assert sheet["C4"].value.date() == date(2026, 8, 11)
    assert sheet["D4"].value == 3900
    assert sheet["L4"].value == "Yes"
    assert sheet["A5"].value == "Replacement savings"
    assert sheet["D5"].value == 3000
    assert sheet["A7"].value == "Planned Withdrawals"
    assert sheet["A9"].value == "August planned expense"
    assert sheet["B9"].value.date() == date(2026, 8, 11)
    assert sheet["D9"].value is True
    assert sheet["E9"].value > 0
    assert sheet["G9"].value == 0
    assert sheet["H9"].value.date() == date(2026, 8, 13)
    assert sheet["A12"].value == "Savings Progress Detail"
    assert sheet["B30"].value == "Replacement savings"
    assert sheet["F30"].value == 1000
    assert sheet["I30"].value > 0
    assert sheet["K30"].value > 0
    assert sheet["M30"].value == 0
    assert sheet["Q30"].value == 3000
    assert sheet["S30"].value == "Aug 13"
    assert sheet.column_dimensions["S"].hidden is True
    assert dashboard["A12"].value == "Active Savings Goal"
    assert dashboard["B12"].value == "Replacement savings"
    assert dashboard["A13"].value == "Active Savings Target"
    assert dashboard["B13"].value == 3000
    assert dashboard["A18"].value == "Goal Feasible"
    assert dashboard["A19"].value == "Snowball Currently Reduced"
    workbook.close()


def test_current_plan_workbook_preserves_disabled_deadline_priority(tmp_path):
    config = Config().load("config.json")
    forecast_config = deepcopy(config)
    periods = CalendarEngine(config.settings).generate(date(2026, 12, 31))
    summaries = BudgetEngine(config).build_plan(periods)
    forecast = ForecastEngine(forecast_config).forecast()
    workbook_path = tmp_path / "current_plan.xlsx"

    ExcelWriter(workbook_path).write(summaries, forecast)

    workbook = load_workbook(workbook_path, data_only=True)
    pay_periods = workbook["Pay Period Summaries"]
    savings = workbook["Savings Progress"]

    assert value_for_header(pay_periods, "Savings Goal", 2) is None
    assert money(value_for_header(pay_periods, "Bills Paid", 2)) == Decimal("766.00")
    assert money(value_for_header(pay_periods, "Savings Deposit", 2)) == Decimal("215.00")
    assert money(value_for_header(pay_periods, "Snowball Redirected To Savings", 2)) == Decimal("0.00")
    assert money(value_for_header(pay_periods, "Personal Expense Reduction", 2)) == Decimal("0.00")
    assert money(value_for_header(pay_periods, "Snowball Payment", 2)) == Decimal("215.00")
    assert money(value_for_header(pay_periods, "Savings Deposit", 3)) == Decimal("568.50")
    assert money(value_for_header(pay_periods, "Snowball Payment", 3)) == Decimal("568.50")
    assert savings["A1"].value == "Pay Date"
    assert money(savings["B2"].value) == Decimal("215.00")
    workbook.close()


def test_enabled_current_plan_workbook_reports_temporary_deadline_goal(tmp_path):
    config = Config().load("config.json")
    config.savings_plan.deadline_priority_enabled = True
    forecast_config = deepcopy(config)
    periods = CalendarEngine(config.settings).generate(date(2026, 12, 31))
    summaries = BudgetEngine(config).build_plan(periods)
    forecast = ForecastEngine(forecast_config).forecast()
    workbook_path = tmp_path / "enabled_plan.xlsx"

    ExcelWriter(workbook_path).write(summaries, forecast)

    workbook = load_workbook(workbook_path, data_only=True)
    pay_periods = workbook["Pay Period Summaries"]
    savings = workbook["Savings Progress"]

    assert money(value_for_header(pay_periods, "Bills Paid", 2)) == Decimal("766.00")
    assert money(value_for_header(pay_periods, "Savings Deposit", 2)) == Decimal("430.00")
    assert money(value_for_header(pay_periods, "Snowball Redirected To Savings", 2)) == Decimal("215.00")
    assert money(value_for_header(pay_periods, "Personal Expense Reduction", 2)) == Decimal("0.00")
    assert money(value_for_header(pay_periods, "Snowball Payment", 2)) == Decimal("0.00")
    assert money(value_for_header(pay_periods, "Savings Deposit", 3)) == Decimal("470.00")
    assert money(value_for_header(pay_periods, "Snowball Payment", 3)) == Decimal("667.00")
    assert value_for_header(pay_periods, "Active Debt Total", 4) > 0
    assert savings["A4"].value == "August withdrawal"
    assert money(savings["I4"].value) == Decimal("2400.00")
    assert money(savings["J4"].value) == Decimal("0.00")
    assert savings["L4"].value == "Yes"
    assert money(savings["E8"].value) == Decimal("2400.00")
    assert money(savings["G8"].value) == Decimal("0.00")
    workbook.close()
