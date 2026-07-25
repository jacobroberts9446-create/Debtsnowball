import csv
import json
import sqlite3
from contextlib import closing
from concurrent.futures import ThreadPoolExecutor
from copy import deepcopy
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from app.budget_engine import BudgetEngine
from app.calendar_engine import CalendarEngine
from app.config import Config
from app.forecast_engine import ForecastEngine
from app.history import (
    PlanHistoryService,
    build_warnings,
    config_fingerprint,
    explain_forecast_period,
    forecast_fingerprint,
    normalized_config_snapshot,
    portable_content_fingerprint,
)
from app.excel_writer import ExcelWriter
from app.models import (
    ActualEntryType,
    AllocationReasonCode,
    DataWarningSeverity,
)
from app.money import money


def short_plan(config):
    forecast = ForecastEngine(config, max_years=1).forecast()
    periods = CalendarEngine(config.settings).generate(forecast.forecast_end_date)
    summaries = BudgetEngine(config).build_plan(periods)
    return summaries, forecast


def history_plan_with_forecast(service, name="Correction Plan"):
    """Create one saved plan with forecast periods for actual-entry tests."""
    config = Config().load("config.json")
    summaries, forecast = short_plan(config)
    plan = service.create_plan(name, config)
    version = service.list_plan_versions(plan.id)[0]
    service.generate_and_save_forecast(
        version.id,
        forecast,
        starting_savings=config.settings.starting_savings,
        pay_period_summaries=summaries,
        starting_debts=config.debts,
    )
    return config, plan


def exported_plan_payload(tmp_path, *, name="Tamper Plan"):
    service = PlanHistoryService(tmp_path / f"{name.replace(' ', '_')}.sqlite")
    config = Config().load("config.json")
    summaries, forecast = short_plan(config)
    plan = service.create_plan(name, config)
    version = service.list_plan_versions(plan.id)[0]
    service.generate_and_save_forecast(
        version.id,
        forecast,
        starting_savings=config.settings.starting_savings,
        pay_period_summaries=summaries,
        starting_debts=config.debts,
        warnings=build_warnings(forecast),
    )
    path = service.export_plan_json(plan.id, tmp_path / f"{name.replace(' ', '_')}.json")
    return json.loads(path.read_text(encoding="utf-8"))


def write_payload(tmp_path, payload, name="payload.json"):
    path = tmp_path / name
    path.write_text(json.dumps(payload), encoding="utf-8")
    return path


def test_configuration_fingerprint_is_stable_and_money_normalized():
    first = {
        "budget": {
            "paycheck": 215,
            "starting_savings": Decimal("-0.00"),
            "first_paycheck": "2026-07-17",
        }
    }
    second = {
        "budget": {
            "first_paycheck": "2026-07-17",
            "starting_savings": "0.00",
            "paycheck": "215.00",
        }
    }
    changed = deepcopy(second)
    changed["budget"]["paycheck"] = "216.00"

    assert normalized_config_snapshot(first)["budget"]["paycheck"] == "215.00"
    assert config_fingerprint(first) == config_fingerprint(second)
    assert config_fingerprint(second) != config_fingerprint(changed)


def test_create_plan_versions_prevent_duplicates_force_and_restore(tmp_path):
    service = PlanHistoryService(tmp_path / "history.sqlite")
    config = Config().load("config.json")

    plan = service.create_plan("Current Plan", config, description="Live config")
    versions = service.list_plan_versions(plan.id)
    assert versions[0].version_number == 1
    assert versions[0].active is True

    duplicate = service.save_plan_version(plan.id, config)
    assert duplicate.id == versions[0].id
    assert len(service.list_plan_versions(plan.id)) == 1

    forced = service.save_plan_version(plan.id, config, force=True, change_note="Forced")
    assert forced.version_number == 2

    changed_config = deepcopy(config)
    changed_config.settings.paycheck = Decimal("2300.00")
    changed = service.save_plan_version(plan.id, changed_config, change_note="Changed")
    assert changed.version_number == 3
    assert service.get_plan_version(versions[0].id).config_snapshot == versions[0].config_snapshot

    restored = service.restore_plan_version(versions[0].id)
    assert restored.version_number == 4
    service.archive_plan(plan.id)
    assert service.list_plans() == []
    assert service.list_plans(include_archived=True)[0].archived is True


def test_save_forecast_snapshot_periods_debt_and_savings_rows_are_exact(tmp_path):
    service = PlanHistoryService(tmp_path / "history.sqlite")
    config = Config().load("config.json")
    summaries, forecast = short_plan(config)
    plan = service.create_plan("Forecast Plan", config)
    version = service.list_plan_versions(plan.id)[0]

    snapshot = service.generate_and_save_forecast(
        version.id,
        forecast,
        starting_savings=config.settings.starting_savings,
        pay_period_summaries=summaries,
        starting_debts=config.debts,
    )

    assert snapshot.forecast_fingerprint == forecast_fingerprint(forecast)
    with closing(sqlite3.connect(tmp_path / "history.sqlite")) as conn:
        stored_interest = conn.execute(
            "SELECT COALESCE(SUM(interest_charged), 0) FROM debt_snapshots",
        ).fetchone()[0]
        period = conn.execute(
            """
            SELECT income, savings_deposit, snowball_payment
            FROM forecast_periods
            WHERE forecast_snapshot_id = ?
            ORDER BY sequence_number
            LIMIT 1
            """,
            (snapshot.id,),
        ).fetchone()
        debt_rows = conn.execute(
            """
            SELECT debt_name, ending_balance, payoff_order, scheduled_minimum_payment
            FROM debt_snapshots
            WHERE forecast_period_id IN (
                SELECT id FROM forecast_periods WHERE forecast_snapshot_id = ?
            )
            ORDER BY forecast_period_id, payoff_order
            """,
            (snapshot.id,),
        ).fetchall()
        savings = conn.execute("SELECT ending_savings FROM savings_snapshots").fetchone()[0]

    assert period == (223400, 21500, 21500)
    assert snapshot.total_projected_interest == money(
        Decimal(stored_interest) / Decimal("100")
    )
    assert isinstance(period[0], int)
    assert len(debt_rows) < len(config.debts) * len(forecast.periods)
    assert len(debt_rows[: len(config.debts)]) == len(config.debts)
    assert "total_debt_summary" not in {row[0] for row in debt_rows}
    assert debt_rows[0][2] == 1
    assert isinstance(debt_rows[0][1], int)
    assert money(
        sum(Decimal(row[1]) for row in debt_rows[: len(config.debts)])
        / Decimal("100")
    ) == forecast.periods[0].total_debt_balance
    assert money(Decimal(savings) / Decimal("100")) == forecast.periods[0].savings_balance


def test_forecast_fingerprint_is_deterministic_for_regenerated_forecast(tmp_path):
    config = Config().load("config.json")
    first = ForecastEngine(config, max_years=1).forecast()
    second = ForecastEngine(config, max_years=1).forecast()

    assert forecast_fingerprint(first) == forecast_fingerprint(second)


def test_actual_entries_reverse_and_compare_without_moralizing(tmp_path):
    service = PlanHistoryService(tmp_path / "history.sqlite")
    config = Config().load("config.json")
    summaries, forecast = short_plan(config)
    plan = service.create_plan("Actual Plan", config)
    version = service.list_plan_versions(plan.id)[0]
    service.generate_and_save_forecast(
        version.id,
        forecast,
        starting_savings=config.settings.starting_savings,
        pay_period_summaries=summaries,
    )

    income = service.add_actual_entry(
        plan.id,
        date(2026, 7, 17),
        ActualEntryType.INCOME_RECEIVED,
        Decimal("2234.00"),
    )
    service.add_actual_entry(
        plan.id,
        date(2026, 7, 17),
        ActualEntryType.DEBT_PAYMENT,
        Decimal("1253.00"),
    )
    reversal = service.reverse_actual_entry(income.id)
    comparison = service.compare_forecast_to_actual(plan.id)
    entries = service.list_actual_entries(plan.id)

    assert reversal.amount == Decimal("-2234.00")
    assert reversal.corrected_entry_id == income.id
    assert reversal.source == "correction"
    assert income in entries
    assert reversal in entries
    assert sum(
        entry.amount
        for entry in entries
        if entry.entry_type == ActualEntryType.INCOME_RECEIVED
    ) == Decimal("0.00")
    assert comparison.status in {
        "Ahead of plan",
        "On track",
        "Needs review",
        "Insufficient actual data",
    }
    assert "bad" not in comparison.status.casefold()
    assert "failed" not in comparison.status.casefold()

    with pytest.raises(ValueError, match="already been reversed"):
        service.reverse_actual_entry(income.id, plan_id=plan.id)
    with pytest.raises(ValueError, match="reversal entry"):
        service.reverse_actual_entry(reversal.id, plan_id=plan.id)

    other_plan = service.create_plan("Other Actual Plan", config)
    with pytest.raises(ValueError, match="selected plan"):
        service.reverse_actual_entry(income.id, plan_id=other_plan.id)
    with pytest.raises(ValueError, match="recorded activity was not found"):
        service.reverse_actual_entry(999999, plan_id=plan.id)

    assert service.list_actual_entries(plan.id) == entries


@pytest.mark.parametrize(
    ("adjustment", "expected_remaining"),
    [
        (Decimal("12.34"), Decimal("12.34")),
        (Decimal("-4.56"), Decimal("-4.56")),
    ],
)
def test_aggregate_signed_adjustments_change_remaining_cash(
    tmp_path,
    adjustment,
    expected_remaining,
):
    service = PlanHistoryService(tmp_path / "adjustment-summary.sqlite")
    config = Config().load("config.json")
    plan = service.create_plan("Adjustment Summary", config)
    service.add_actual_entry(
        plan.id,
        config.settings.first_paycheck,
        ActualEntryType.ADJUSTMENT,
        adjustment,
    )

    comparison = service.compare_forecast_to_actual(plan.id)

    assert comparison.actual_remaining_cash == expected_remaining


def test_aggregate_savings_withdrawal_reduces_net_savings_and_releases_cash(tmp_path):
    service = PlanHistoryService(tmp_path / "withdrawal-summary.sqlite")
    config = Config().load("config.json")
    plan = service.create_plan("Withdrawal Summary", config)
    service.add_actual_entry(
        plan.id,
        config.settings.first_paycheck,
        ActualEntryType.SAVINGS_DEPOSIT,
        Decimal("100.05"),
    )
    service.add_actual_entry(
        plan.id,
        config.settings.first_paycheck,
        ActualEntryType.SAVINGS_WITHDRAWAL,
        Decimal("40.02"),
    )

    comparison = service.compare_forecast_to_actual(plan.id)

    assert comparison.actual_savings == Decimal("60.03")
    assert comparison.actual_remaining_cash == Decimal("-60.03")


def test_aggregate_mixed_activity_uses_exact_decimal_cash_flow(tmp_path):
    service = PlanHistoryService(tmp_path / "mixed-summary.sqlite")
    config = Config().load("config.json")
    plan = service.create_plan("Mixed Summary", config)
    activity = [
        (ActualEntryType.INCOME_RECEIVED, Decimal("2000.10")),
        (ActualEntryType.BILL_PAID, Decimal("500.05")),
        (ActualEntryType.DEBT_PAYMENT, Decimal("300.01")),
        (ActualEntryType.SAVINGS_DEPOSIT, Decimal("200.02")),
        (ActualEntryType.SAVINGS_WITHDRAWAL, Decimal("50.03")),
        (ActualEntryType.PERSONAL_SPENDING, Decimal("100.04")),
        (ActualEntryType.ADJUSTMENT, Decimal("10.05")),
        (ActualEntryType.ADJUSTMENT, Decimal("-2.06")),
    ]
    for entry_type, amount in activity:
        service.add_actual_entry(
            plan.id,
            config.settings.first_paycheck,
            entry_type,
            amount,
        )

    comparison = service.compare_forecast_to_actual(plan.id)

    assert comparison.actual_income == Decimal("2000.10")
    assert comparison.actual_bills == Decimal("500.05")
    assert comparison.actual_debt_payments == Decimal("300.01")
    assert comparison.actual_savings == Decimal("149.99")
    assert comparison.actual_personal_spending == Decimal("100.04")
    assert comparison.actual_remaining_cash == Decimal("958.00")


def test_atomic_actual_correction_preserves_rows_periods_and_totals(tmp_path):
    service = PlanHistoryService(tmp_path / "correction.sqlite")
    _, plan = history_plan_with_forecast(service)
    original = service.add_actual_entry(
        plan.id,
        date(2026, 7, 20),
        ActualEntryType.INCOME_RECEIVED,
        Decimal("2000.00"),
        category="Income",
        note="Original",
    )

    result = service.correct_actual_entry(
        plan.id,
        original.id,
        entry_date=date(2026, 8, 1),
        amount=Decimal("2100.00"),
        category="Income",
        debt_identifier=None,
        note="Corrected paycheck",
    )

    assert result.original == original
    assert result.reversal.amount == Decimal("-2000.00")
    assert result.reversal.source == "correction"
    assert result.reversal.corrected_entry_id == original.id
    assert result.replacement.amount == Decimal("2100.00")
    assert result.replacement.entry_date == date(2026, 8, 1)
    assert result.replacement.note == "Corrected paycheck"
    assert result.replacement.source == "correction_replacement"
    assert result.replacement.corrected_entry_id == original.id
    assert service.compare_forecast_to_actual(plan.id).actual_income == Decimal(
        "2100.00"
    )

    with closing(sqlite3.connect(tmp_path / "correction.sqlite")) as conn:
        rows = conn.execute(
            """
            SELECT id, forecast_period_id
            FROM actual_transactions
            WHERE id IN (?, ?, ?)
            ORDER BY id
            """,
            (original.id, result.reversal.id, result.replacement.id),
        ).fetchall()
    assert rows[0][1] == rows[1][1]
    assert rows[2][1] != rows[0][1]


def test_atomic_correction_supports_associations_and_replacement_chains(tmp_path):
    service = PlanHistoryService(tmp_path / "chains.sqlite")
    _, plan = history_plan_with_forecast(service)
    bill = service.add_actual_entry(
        plan.id,
        date(2026, 7, 20),
        ActualEntryType.BILL_PAID,
        Decimal("80.00"),
        category="Internet",
    )
    first = service.correct_actual_entry(
        plan.id,
        bill.id,
        entry_date=bill.entry_date,
        amount=Decimal("85.00"),
        category="Fiber Internet",
        debt_identifier=None,
        note="Updated bill",
    )
    second = service.correct_actual_entry(
        plan.id,
        first.replacement.id,
        entry_date=date(2026, 7, 21),
        amount=Decimal("90.00"),
        category="Fiber Internet",
        debt_identifier=None,
        note="Final bill",
    )
    debt = service.add_actual_entry(
        plan.id,
        date(2026, 7, 20),
        ActualEntryType.DEBT_PAYMENT,
        Decimal("100.00"),
        category="Debt Payment",
        debt_identifier="Card A",
    )
    debt_result = service.correct_actual_entry(
        plan.id,
        debt.id,
        entry_date=debt.entry_date,
        amount=Decimal("125.00"),
        category="Debt Payment",
        debt_identifier="Card B",
        note="Correct card",
    )

    assert second.reversal.corrected_entry_id == first.replacement.id
    assert second.replacement.corrected_entry_id == first.replacement.id
    assert second.replacement.category == "Fiber Internet"
    assert debt_result.replacement.debt_identifier == "Card B"
    comparison = service.compare_forecast_to_actual(plan.id)
    assert comparison.actual_bills == Decimal("90.00")
    assert comparison.actual_debt_payments == Decimal("125.00")

    final_reversal = service.reverse_actual_entry(
        second.replacement.id,
        plan_id=plan.id,
    )
    assert final_reversal.corrected_entry_id == second.replacement.id
    assert service.compare_forecast_to_actual(plan.id).actual_bills == Decimal(
        "0.00"
    )


@pytest.mark.parametrize("failure_source", ["correction", "correction_replacement"])
def test_atomic_correction_rolls_back_both_rows_on_insert_failure(
    tmp_path,
    failure_source,
):
    database_path = tmp_path / f"rollback-{failure_source}.sqlite"
    service = PlanHistoryService(database_path)
    _, plan = history_plan_with_forecast(service)
    original = service.add_actual_entry(
        plan.id,
        date(2026, 7, 20),
        ActualEntryType.PERSONAL_SPENDING,
        Decimal("25.00"),
        category="Personal Spending",
    )
    with closing(sqlite3.connect(database_path)) as conn:
        conn.execute(
            f"""
            CREATE TRIGGER fail_correction_insert
            BEFORE INSERT ON actual_transactions
            WHEN NEW.source = '{failure_source}'
            BEGIN
                SELECT RAISE(ABORT, 'injected correction failure');
            END
            """
        )

    with pytest.raises(sqlite3.IntegrityError, match="injected correction failure"):
        service.correct_actual_entry(
            plan.id,
            original.id,
            entry_date=original.entry_date,
            amount=Decimal("30.00"),
            category=original.category,
            debt_identifier=None,
            note="Should roll back",
        )

    assert service.list_actual_entries(plan.id) == [original]


def test_atomic_correction_validates_targets_signs_and_associations(tmp_path):
    service = PlanHistoryService(tmp_path / "validation.sqlite")
    _, plan = history_plan_with_forecast(service)
    other_plan = service.create_plan("Other Correction Plan", Config().load("config.json"))
    original = service.add_actual_entry(
        plan.id,
        date(2026, 7, 20),
        ActualEntryType.INCOME_RECEIVED,
        Decimal("100.00"),
        category="Income",
    )

    base = {
        "entry_date": original.entry_date,
        "amount": Decimal("110.00"),
        "category": original.category,
        "debt_identifier": None,
        "note": "",
    }
    with pytest.raises(ValueError, match="selected plan"):
        service.correct_actual_entry(other_plan.id, original.id, **base)
    with pytest.raises(ValueError, match="recorded activity was not found"):
        service.correct_actual_entry(plan.id, 999999, **base)
    with pytest.raises(ValueError, match="greater than"):
        service.correct_actual_entry(
            plan.id,
            original.id,
            **{**base, "amount": Decimal("0.00")},
        )
    with pytest.raises(ValueError, match="debt name"):
        service.correct_actual_entry(
            plan.id,
            original.id,
            **{**base, "debt_identifier": "Card A"},
        )

    adjustment = service.add_actual_entry(
        plan.id,
        date(2026, 7, 20),
        ActualEntryType.ADJUSTMENT,
        Decimal("10.00"),
        category="Adjustment",
    )
    adjustment_base = {
        "entry_date": adjustment.entry_date,
        "category": adjustment.category,
        "debt_identifier": None,
        "note": "",
    }
    negative = service.correct_actual_entry(
        plan.id,
        adjustment.id,
        amount=Decimal("-5.00"),
        **adjustment_base,
    )
    assert negative.replacement.amount == Decimal("-5.00")
    with pytest.raises(ValueError, match="nonzero"):
        service.correct_actual_entry(
            plan.id,
            negative.replacement.id,
            amount=Decimal("0.00"),
            **adjustment_base,
        )


def test_atomic_correction_rejects_reversal_and_superseded_targets(tmp_path):
    service = PlanHistoryService(tmp_path / "eligibility.sqlite")
    _, plan = history_plan_with_forecast(service)
    original = service.add_actual_entry(
        plan.id,
        date(2026, 7, 20),
        ActualEntryType.SAVINGS_DEPOSIT,
        Decimal("50.00"),
        category="Savings",
    )
    result = service.correct_actual_entry(
        plan.id,
        original.id,
        entry_date=original.entry_date,
        amount=Decimal("75.00"),
        category=original.category,
        debt_identifier=None,
        note="Corrected",
    )
    replacement = {
        "entry_date": original.entry_date,
        "amount": Decimal("80.00"),
        "category": original.category,
        "debt_identifier": None,
        "note": "",
    }

    with pytest.raises(ValueError, match="already been superseded"):
        service.correct_actual_entry(plan.id, original.id, **replacement)
    with pytest.raises(ValueError, match="reversal entry"):
        service.correct_actual_entry(plan.id, result.reversal.id, **replacement)


def test_concurrent_atomic_corrections_allow_only_one_winner(tmp_path):
    service = PlanHistoryService(tmp_path / "concurrent.sqlite")
    _, plan = history_plan_with_forecast(service)
    original = service.add_actual_entry(
        plan.id,
        date(2026, 7, 20),
        ActualEntryType.INCOME_RECEIVED,
        Decimal("100.00"),
        category="Income",
    )

    def correct(amount):
        try:
            return service.correct_actual_entry(
                plan.id,
                original.id,
                entry_date=original.entry_date,
                amount=amount,
                category=original.category,
                debt_identifier=None,
                note="Concurrent",
            )
        except ValueError as exc:
            return str(exc)

    with ThreadPoolExecutor(max_workers=2) as executor:
        results = list(
            executor.map(correct, [Decimal("110.00"), Decimal("120.00")])
        )

    assert sum(not isinstance(result, str) for result in results) == 1
    assert sum("already been superseded" in result for result in results if isinstance(result, str)) == 1
    assert len(service.list_actual_entries(plan.id)) == 3


def test_correction_chain_survives_json_export_import(tmp_path):
    service = PlanHistoryService(tmp_path / "export-source.sqlite")
    _, plan = history_plan_with_forecast(service, "Correction Export")
    original = service.add_actual_entry(
        plan.id,
        date(2026, 7, 20),
        ActualEntryType.INCOME_RECEIVED,
        Decimal("100.00"),
        category="Income",
    )
    first = service.correct_actual_entry(
        plan.id,
        original.id,
        entry_date=original.entry_date,
        amount=Decimal("110.00"),
        category="Income",
        debt_identifier=None,
        note="First",
    )
    service.correct_actual_entry(
        plan.id,
        first.replacement.id,
        entry_date=date(2026, 7, 21),
        amount=Decimal("120.00"),
        category="Income",
        debt_identifier=None,
        note="Second",
    )

    export_path = service.export_plan_json(plan.id, tmp_path / "correction.json")
    imported_service = PlanHistoryService(tmp_path / "export-target.sqlite")
    imported_plan = imported_service.import_plan_json(export_path)
    entries = imported_service.list_actual_entries(imported_plan.id)

    assert len(entries) == 5
    assert [entry.source for entry in entries].count("correction") == 2
    assert [entry.source for entry in entries].count("correction_replacement") == 2
    assert (
        imported_service.compare_forecast_to_actual(imported_plan.id).actual_income
        == Decimal("120.00")
    )


@pytest.mark.parametrize(
    ("children", "message"),
    [
        (
            [
                {
                    "id": 2,
                    "entry_type": "income_received",
                    "entry_date": "2026-07-20",
                    "amount": "-100.00",
                    "source": "correction_replacement",
                    "corrected_entry_id": 1,
                    "forecast_period_id": None,
                }
            ],
            "one reversal",
        ),
        (
            [
                {
                    "id": 2,
                    "entry_type": "income_received",
                    "entry_date": "2026-07-20",
                    "amount": "-99.99",
                    "source": "correction",
                    "corrected_entry_id": 1,
                    "forecast_period_id": None,
                }
            ],
            "reversal amount mismatch",
        ),
    ],
)
def test_import_rejects_malformed_correction_relationships(
    tmp_path,
    children,
    message,
):
    service = PlanHistoryService(tmp_path / "invalid-correction.sqlite")
    actual_entries = [
        {
            "id": 1,
            "entry_type": "income_received",
            "entry_date": "2026-07-20",
            "amount": "100.00",
            "source": "interactive",
            "corrected_entry_id": None,
            "forecast_period_id": None,
        },
        *children,
    ]

    with pytest.raises(ValueError, match=message):
        service._importer().validate_import_actuals(
            {"actual_entries": actual_entries, "balance_observations": []},
            set(),
        )


def test_plan_comparison_reports_tradeoffs_and_first_difference(tmp_path):
    service = PlanHistoryService(tmp_path / "history.sqlite")
    config = Config().load("config.json")
    plan = service.create_plan("Compare Plan", config)
    version_one = service.list_plan_versions(plan.id)[0]
    summaries_one, forecast_one = short_plan(config)
    service.generate_and_save_forecast(
        version_one.id,
        forecast_one,
        starting_savings=config.settings.starting_savings,
        pay_period_summaries=summaries_one,
    )

    changed_config = deepcopy(config)
    changed_config.settings.paycheck = Decimal("2500.00")
    version_two = service.save_plan_version(plan.id, changed_config)
    summaries_two, forecast_two = short_plan(changed_config)
    service.generate_and_save_forecast(
        version_two.id,
        forecast_two,
        starting_savings=changed_config.settings.starting_savings,
        pay_period_summaries=summaries_two,
    )

    comparison = service.compare_plan_versions(version_one.id, version_two.id)

    assert comparison.later_version == 2
    assert comparison.first_different_period == date(2026, 7, 17)
    assert comparison.interest_difference <= Decimal("0.00")
    assert "better" not in comparison.explanation.casefold()


def test_plan_comparison_accepts_null_savings_plan_snapshots(tmp_path):
    """Interactive config snapshots with no savings plan remain comparable."""
    service = PlanHistoryService(tmp_path / "history.sqlite")
    config = Config().load("config.json")
    config.savings_plan = None
    summaries, forecast = short_plan(config)

    first = service.save_generated_plan(
        name="Interactive Plan",
        config=config,
        forecast=forecast,
        starting_savings=config.settings.starting_savings,
        starting_debts=config.debts,
        pay_period_summaries=summaries,
    )
    second = service.save_generated_plan(
        name="",
        plan_id=first.plan.id,
        config=config,
        forecast=forecast,
        starting_savings=config.settings.starting_savings,
        starting_debts=config.debts,
        pay_period_summaries=summaries,
        force=True,
    )

    comparison = service.compare_plan_versions(first.version.id, second.version.id)

    assert comparison.deadline_priority_changed is False


def test_export_import_json_and_csv_preserve_money_strings_and_duplicate_protection(tmp_path):
    service = PlanHistoryService(tmp_path / "source.sqlite")
    config = Config().load("config.json")
    summaries, forecast = short_plan(config)
    plan = service.create_plan("Export Plan", config)
    version = service.list_plan_versions(plan.id)[0]
    snapshot = service.generate_and_save_forecast(
        version.id,
        forecast,
        starting_savings=config.settings.starting_savings,
        pay_period_summaries=summaries,
        starting_debts=config.debts,
    )
    service.add_actual_entry(
        plan.id,
        date(2026, 7, 17),
        ActualEntryType.INCOME_RECEIVED,
        Decimal("2234.00"),
    )
    service.add_balance_observation(
        plan.id,
        date(2026, 7, 18),
        ActualEntryType.SAVINGS_BALANCE_OBSERVATION,
        Decimal("1715.00"),
    )

    json_path = service.export_plan_json(plan.id, tmp_path / "plan.json")
    csv_path = service.export_forecast_periods_csv(snapshot.id, tmp_path / "periods.csv")
    csv_bundle = service.export_csv_bundle(plan.id, tmp_path / "csv")
    payload = json.loads(json_path.read_text(encoding="utf-8"))
    with csv_path.open(encoding="utf-8") as csv_file:
        csv_rows = list(csv.DictReader(csv_file))

    assert payload["versions"][0]["config_snapshot"].count("2400.00") > 0
    assert csv_rows[0]["savings"] == "215.00"
    assert set(csv_bundle) == {
        "plan_versions",
        "forecast_periods",
        "debt_history",
        "savings_history",
        "actual_transactions",
        "forecast_vs_actual",
        "warnings",
    }
    assert csv_bundle["debt_history"].exists()

    imported_service = PlanHistoryService(tmp_path / "target.sqlite")
    imported = imported_service.import_plan_json(json_path)
    assert imported.name == "Export Plan"
    imported_snapshot = imported_service._latest_snapshot_for_version(
        imported.current_version_id
    )
    assert len(imported_service._exporter().export_forecast_periods(imported.id)) == len(
        forecast.periods
    )
    assert len(imported_service._exporter().export_debt_snapshots(imported.id)) == len(
        payload["debt_snapshots"]
    )
    assert imported_service.get_forecast_snapshot(
        imported_snapshot.id
    ).total_projected_interest == snapshot.total_projected_interest
    assert imported_service.list_actual_entries(imported.id)[0].amount == Decimal(
        "2234.00"
    )
    assert imported_service.list_balance_observations(imported.id)[0].balance == Decimal(
        "1715.00"
    )
    with pytest.raises(ValueError, match="already been imported"):
        imported_service.import_plan_json(json_path)

    bad_path = tmp_path / "bad.json"
    bad_path.write_text('{"format_version": 999}', encoding="utf-8")
    with pytest.raises(ValueError, match="unsupported"):
        imported_service.import_plan_json(bad_path, new_name="Bad")
    assert len(imported_service.list_plans()) == 1


def test_optional_history_excel_report_uses_requested_history_sheets(tmp_path):
    service = PlanHistoryService(tmp_path / "history.sqlite")
    config = Config().load("config.json")
    summaries, forecast = short_plan(config)
    plan = service.create_plan("Workbook Plan", config)
    version = service.list_plan_versions(plan.id)[0]
    service.generate_and_save_forecast(
        version.id,
        forecast,
        starting_savings=config.settings.starting_savings,
        pay_period_summaries=summaries,
        starting_debts=config.debts,
    )
    service.add_actual_entry(
        plan.id,
        date(2026, 7, 17),
        ActualEntryType.INCOME_RECEIVED,
        Decimal("2234.00"),
    )
    actual = service.compare_forecast_to_actual(plan.id)

    details = service.history_report_rows(plan.id)
    path = ExcelWriter(tmp_path / "history.xlsx").write_history_report(
        plan,
        service.list_plan_versions(plan.id),
        actual_comparison=actual,
        forecast_snapshots=details["forecast_snapshots"],
        actual_periods=service.compare_forecast_to_actual_periods(plan.id),
        debt_history=details["debt_history"],
        savings_history=details["savings_history"],
        warnings=details["warnings"],
    )

    workbook = load_workbook(path)
    assert workbook.sheetnames == [
        "Plan History",
        "Forecast vs Actual",
        "Period Details",
        "Debt History",
        "Savings History",
        "History Warnings",
    ]
    assert workbook["Forecast vs Actual"]["B3"].value > 8936
    assert workbook["Forecast vs Actual"]["C3"].value == 2234
    assert workbook["Plan History"]["E2"].value == "Forecast Fingerprint"
    assert isinstance(workbook["Plan History"]["G3"].value, int | float)
    assert workbook["Debt History"]["A1"].value == "id"
    assert workbook["Savings History"]["A1"].value == "id"
    assert isinstance(workbook["Forecast vs Actual"]["B3"].value, int | float)
    workbook.close()


def test_actual_entries_and_observations_match_forecast_period_windows(tmp_path):
    service = PlanHistoryService(tmp_path / "history.sqlite")
    config = Config().load("config.json")
    summaries, forecast = short_plan(config)
    plan = service.create_plan("Match Plan", config)
    version = service.list_plan_versions(plan.id)[0]
    service.generate_and_save_forecast(
        version.id,
        forecast,
        starting_savings=config.settings.starting_savings,
        pay_period_summaries=summaries,
        starting_debts=config.debts,
    )

    matched = service.add_actual_entry(
        plan.id,
        date(2026, 7, 20),
        ActualEntryType.BILL_PAID,
        Decimal("100.00"),
    )
    observation = service.add_balance_observation(
        plan.id,
        date(2026, 7, 20),
        ActualEntryType.DEBT_BALANCE_OBSERVATION,
        forecast.periods[0].total_debt_balance - Decimal("5.00"),
    )
    unmatched = service.add_actual_entry(
        plan.id,
        date(2026, 1, 1),
        ActualEntryType.BILL_PAID,
        Decimal("10.00"),
    )
    comparisons = service.compare_forecast_to_actual_periods(plan.id)

    with closing(sqlite3.connect(tmp_path / "history.sqlite")) as conn:
        rows = conn.execute(
            """
            SELECT id, forecast_period_id, match_method
            FROM actual_transactions
            ORDER BY id
            """
        ).fetchall()

    assert matched.id == rows[0][0]
    assert rows[0][1] == comparisons[0].forecast_period_id
    assert rows[0][2] == "date_window"
    assert observation.forecast_period_id == comparisons[0].forecast_period_id
    assert unmatched.id == rows[1][0]
    assert rows[1][1] is None
    assert comparisons[0].debt_balance_variance == Decimal("-5.00")


def test_progress_is_rematched_to_new_version_periods_without_mutating_history(tmp_path):
    service = PlanHistoryService(tmp_path / "version-progress.sqlite")
    config = Config().load("config.json")
    summaries, forecast = short_plan(config)
    plan = service.create_plan("Version Progress", config)
    version_one = service.list_plan_versions(plan.id)[0]
    service.generate_and_save_forecast(
        version_one.id,
        forecast,
        starting_savings=config.settings.starting_savings,
        pay_period_summaries=summaries,
        starting_debts=config.debts,
    )
    activity_date = date(2026, 8, 1)
    service.add_actual_entry(
        plan.id,
        activity_date,
        ActualEntryType.INCOME_RECEIVED,
        Decimal("100.00"),
    )
    service.add_actual_entry(
        plan.id,
        activity_date,
        ActualEntryType.BILL_PAID,
        Decimal("25.00"),
    )
    service.add_actual_entry(
        plan.id,
        date(2026, 8, 10),
        ActualEntryType.BILL_PAID,
        Decimal("15.00"),
    )
    observation = service.add_balance_observation(
        plan.id,
        activity_date,
        ActualEntryType.DEBT_BALANCE_OBSERVATION,
        forecast.periods[1].total_debt_balance,
    )
    original_period_id = observation.forecast_period_id

    changed = deepcopy(config)
    changed.settings.first_paycheck = date(2026, 7, 31)
    changed_summaries, changed_forecast = short_plan(changed)
    version_two = service.save_plan_version(plan.id, changed)
    service.generate_and_save_forecast(
        version_two.id,
        changed_forecast,
        starting_savings=changed.settings.starting_savings,
        pay_period_summaries=changed_summaries,
        starting_debts=changed.debts,
    )

    comparisons = service.compare_forecast_to_actual_periods(plan.id)
    first = comparisons[0]
    with closing(sqlite3.connect(tmp_path / "version-progress.sqlite")) as conn:
        actual_rows = conn.execute(
            """
            SELECT forecast_period_id
            FROM actual_transactions
            WHERE plan_id = ?
            ORDER BY id
            """,
            (plan.id,),
        ).fetchall()
        observation_row = conn.execute(
            """
            SELECT forecast_period_id
            FROM balance_observations
            WHERE id = ?
            """,
            (observation.id,),
        ).fetchone()
        foreign_key_errors = conn.execute("PRAGMA foreign_key_check").fetchall()

    assert first.pay_date == date(2026, 7, 31)
    assert first.actual_income == Decimal("100.00")
    assert first.actual_bills == Decimal("40.00")
    assert first.debt_balance_variance == money(
        forecast.periods[1].total_debt_balance
        - changed_forecast.periods[0].total_debt_balance
    )
    assert len(service.list_actual_entries(plan.id)) == 3
    assert len(service.list_balance_observations(plan.id)) == 1
    assert {row[0] for row in actual_rows} == {original_period_id}
    assert observation_row == (original_period_id,)
    assert first.forecast_period_id != original_period_id
    assert foreign_key_errors == []


def test_progress_survives_multiple_versions_and_restored_version(tmp_path):
    service = PlanHistoryService(tmp_path / "restore-progress.sqlite")
    config = Config().load("config.json")
    summaries, forecast = short_plan(config)
    plan = service.create_plan("Restore Progress", config)
    version_one = service.list_plan_versions(plan.id)[0]
    service.generate_and_save_forecast(
        version_one.id,
        forecast,
        starting_savings=config.settings.starting_savings,
        pay_period_summaries=summaries,
        starting_debts=config.debts,
    )
    pay_date = forecast.periods[0].paycheck_date
    service.add_actual_entry(
        plan.id,
        pay_date,
        ActualEntryType.DEBT_PAYMENT,
        Decimal("75.00"),
    )

    version_two = service.save_plan_version(plan.id, config, force=True)
    service.generate_and_save_forecast(
        version_two.id,
        forecast,
        starting_savings=config.settings.starting_savings,
        pay_period_summaries=summaries,
        starting_debts=config.debts,
    )
    assert service.compare_forecast_to_actual_periods(plan.id)[
        0
    ].actual_debt_payments == Decimal("75.00")

    restored = service.restore_plan_version(version_one.id)
    restored_comparison = service.compare_forecast_to_actual_periods(plan.id)

    assert restored.version_number == 3
    assert restored_comparison[0].actual_debt_payments == Decimal("75.00")
    assert len(service.list_actual_entries(plan.id)) == 1
    with pytest.raises(ValueError, match="does not have a saved forecast"):
        service._latest_snapshot_for_version(restored.id)


def test_period_comparison_reports_complete_on_track_actuals(tmp_path):
    service = PlanHistoryService(tmp_path / "history.sqlite")
    config = Config().load("config.json")
    summaries, forecast = short_plan(config)
    plan = service.create_plan("Complete Actual Plan", config)
    version = service.list_plan_versions(plan.id)[0]
    service.generate_and_save_forecast(
        version.id,
        forecast,
        starting_savings=config.settings.starting_savings,
        pay_period_summaries=summaries,
        starting_debts=config.debts,
    )
    baseline_period = service.compare_forecast_to_actual_periods(plan.id)[0]
    assert baseline_period.status == "No actual activity recorded"
    assert baseline_period.data_completeness == "none"

    entries = [
        (ActualEntryType.INCOME_RECEIVED, baseline_period.planned_income),
        (
            ActualEntryType.BILL_PAID,
            baseline_period.planned_bills - baseline_period.planned_personal_spending,
        ),
        (
            ActualEntryType.DEBT_PAYMENT,
            baseline_period.planned_debt_minimums + baseline_period.planned_snowball,
        ),
        (ActualEntryType.SAVINGS_DEPOSIT, baseline_period.planned_savings_deposit),
        (ActualEntryType.PERSONAL_SPENDING, baseline_period.planned_personal_spending),
    ]
    for entry_type, amount in entries:
        service.add_actual_entry(plan.id, baseline_period.pay_date, entry_type, amount)

    complete_period = service.compare_forecast_to_actual_periods(plan.id)[0]

    assert complete_period.data_completeness == "complete"
    assert complete_period.status == "On track"
    assert complete_period.actual_remaining_cash == baseline_period.planned_remaining_cash
    assert complete_period.interpretation == (
        "Recorded activity is close to the forecast for this period."
    )


def test_period_comparison_includes_withdrawals_and_signed_adjustments(tmp_path):
    service = PlanHistoryService(tmp_path / "period-cash-flow.sqlite")
    config, plan = history_plan_with_forecast(service, "Period Cash Flow")
    activity = [
        (ActualEntryType.INCOME_RECEIVED, Decimal("1000.00")),
        (ActualEntryType.BILL_PAID, Decimal("100.00")),
        (ActualEntryType.DEBT_PAYMENT, Decimal("200.00")),
        (ActualEntryType.SAVINGS_DEPOSIT, Decimal("300.00")),
        (ActualEntryType.SAVINGS_WITHDRAWAL, Decimal("50.00")),
        (ActualEntryType.PERSONAL_SPENDING, Decimal("100.00")),
        (ActualEntryType.ADJUSTMENT, Decimal("25.25")),
        (ActualEntryType.ADJUSTMENT, Decimal("-5.05")),
    ]
    for entry_type, amount in activity:
        service.add_actual_entry(
            plan.id,
            config.settings.first_paycheck,
            entry_type,
            amount,
        )

    comparison = service.compare_forecast_to_actual_periods(plan.id)[0]

    assert comparison.actual_income == Decimal("1000.00")
    assert comparison.actual_bills == Decimal("100.00")
    assert comparison.actual_debt_payments == Decimal("200.00")
    assert comparison.actual_savings_deposit == Decimal("300.00")
    assert comparison.actual_savings_withdrawal == Decimal("50.00")
    assert comparison.actual_personal_spending == Decimal("100.00")
    assert comparison.actual_remaining_cash == Decimal("370.20")


def test_persisted_forecast_detail_reconciles_exactly(tmp_path):
    service = PlanHistoryService(tmp_path / "history.sqlite")
    config = Config().load("config.json")
    summaries, forecast = short_plan(config)
    plan = service.create_plan("Reconcile Plan", config)
    version = service.list_plan_versions(plan.id)[0]
    snapshot = service.generate_and_save_forecast(
        version.id,
        forecast,
        starting_savings=config.settings.starting_savings,
        pay_period_summaries=summaries,
        starting_debts=config.debts,
    )

    with closing(sqlite3.connect(tmp_path / "history.sqlite")) as conn:
        periods = conn.execute(
            """
            SELECT id, income, fixed_expenses, minimum_debt_payments,
                   snowball_payment, savings_deposit, savings_withdrawal,
                   checking_remaining
            FROM forecast_periods
            WHERE forecast_snapshot_id = ?
            """,
            (snapshot.id,),
        ).fetchall()
        debts = conn.execute(
            """
            SELECT forecast_period_id, starting_balance, interest_charged,
                   total_payment, ending_balance
            FROM debt_snapshots
            """
        ).fetchall()
        savings = conn.execute(
            """
            SELECT starting_savings, total_deposit, withdrawal, ending_savings
            FROM savings_snapshots
            """
        ).fetchall()

    for row in periods:
        (
            _period_id,
            income,
            fixed_expenses,
            minimums,
            snowball,
            savings_deposit,
            savings_withdrawal,
            checking,
        ) = row
        assert (
            income
            - fixed_expenses
            - minimums
            - snowball
            - savings_deposit
            + savings_withdrawal
            - checking
        ) == 0

    for _period_id, starting, interest, payment, ending in debts:
        assert starting + interest - payment == ending

    for starting, deposit, withdrawal, ending in savings:
        assert starting + deposit - withdrawal == ending


def test_history_service_safe_delete_exports_and_removes_children(tmp_path):
    service = PlanHistoryService(tmp_path / "history.sqlite")
    config = Config().load("config.json")
    summaries, forecast = short_plan(config)
    plan = service.create_plan("Delete Plan", config)
    version = service.list_plan_versions(plan.id)[0]
    service.generate_and_save_forecast(
        version.id,
        forecast,
        starting_savings=config.settings.starting_savings,
        pay_period_summaries=summaries,
        starting_debts=config.debts,
    )
    service.add_actual_entry(
        plan.id,
        date(2026, 7, 17),
        ActualEntryType.INCOME_RECEIVED,
        Decimal("2234.00"),
    )

    with pytest.raises(ValueError, match="archive"):
        service.delete_plan_permanently(plan.id, confirmation_name="Delete Plan")

    service.archive_plan(plan.id)
    with pytest.raises(ValueError, match="confirmation"):
        service.delete_plan_permanently(plan.id, confirmation_name="Wrong")

    export_path = tmp_path / "delete-plan.json"
    service.delete_plan_permanently(
        plan.id,
        confirmation_name="Delete Plan",
        export_path=export_path,
    )

    assert export_path.exists()
    with pytest.raises(ValueError, match="was not found"):
        service.get_plan(plan.id)
    with closing(sqlite3.connect(tmp_path / "history.sqlite")) as conn:
        assert conn.execute("SELECT COUNT(*) FROM plan_versions").fetchone()[0] == 0
        assert conn.execute("SELECT COUNT(*) FROM forecast_periods").fetchone()[0] == 0
        assert conn.execute("SELECT COUNT(*) FROM debt_snapshots").fetchone()[0] == 0
        assert conn.execute("SELECT COUNT(*) FROM actual_transactions").fetchone()[0] == 0


def test_history_summary_assumption_differences_and_error_paths(tmp_path):
    service = PlanHistoryService(tmp_path / "history.sqlite")
    config = Config().load("config.json")
    summaries, forecast = short_plan(config)
    plan = service.create_plan("Assumption Plan", config)
    version_one = service.list_plan_versions(plan.id)[0]
    service.generate_and_save_forecast(
        version_one.id,
        forecast,
        starting_savings=config.settings.starting_savings,
        pay_period_summaries=summaries,
        starting_debts=config.debts,
    )

    changed = deepcopy(config)
    changed.bills = changed.bills[:-1]
    changed.debts[0].apr = Decimal("30.00")
    changed.debts[0].snowball_order = 6
    changed.debts[-1].snowball_order = 1
    changed.savings_plan.deadline_priority_enabled = (
        not changed.savings_plan.deadline_priority_enabled
    )
    version_two = service.save_plan_version(plan.id, changed)
    changed_summaries, changed_forecast = short_plan(changed)
    service.generate_and_save_forecast(
        version_two.id,
        changed_forecast,
        starting_savings=changed.settings.starting_savings,
        pay_period_summaries=changed_summaries,
        starting_debts=changed.debts,
    )

    summary = service.plan_summary(version_one.id)
    differences = service.compare_assumptions(version_one.id, version_two.id)
    comparison = service.compare_plan_versions(version_one.id, version_two.id)

    assert "starting debt" in summary
    assert {difference.category for difference in differences} >= {
        "Bills",
        "Debts",
        "Savings Plan",
    }
    assert comparison.payoff_order_changed is True
    assert "better" not in comparison.explanation.casefold()
    with pytest.raises(ValueError, match="plan 999999"):
        service.compare_forecast_to_actual_periods(999999)
    with pytest.raises(ValueError, match="balance observation type"):
        service.add_balance_observation(
            plan.id,
            date(2026, 7, 17),
            ActualEntryType.INCOME_RECEIVED,
            Decimal("1.00"),
        )
    with pytest.raises(ValueError, match="balance observation"):
        service.get_balance_observation(999999)
    with pytest.raises(ValueError, match="actual entry"):
        service.get_actual_entry(999999)


def test_import_validation_rejects_nested_fingerprint_mismatch_and_rolls_back(tmp_path):
    service = PlanHistoryService(tmp_path / "source.sqlite")
    config = Config().load("config.json")
    summaries, forecast = short_plan(config)
    plan = service.create_plan("Invalid Import Plan", config)
    version = service.list_plan_versions(plan.id)[0]
    service.generate_and_save_forecast(
        version.id,
        forecast,
        starting_savings=config.settings.starting_savings,
        pay_period_summaries=summaries,
        starting_debts=config.debts,
    )
    export_path = service.export_plan_json(plan.id, tmp_path / "plan.json")
    payload = json.loads(export_path.read_text(encoding="utf-8"))
    payload["versions"][0]["config_fingerprint"] = "bad"
    bad_path = tmp_path / "bad-fingerprint.json"
    bad_path.write_text(json.dumps(payload), encoding="utf-8")

    imported_service = PlanHistoryService(tmp_path / "target.sqlite")
    with pytest.raises(ValueError, match="fingerprint"):
        imported_service.import_plan_json(bad_path)

    assert imported_service.list_plans() == []


def test_portable_content_fingerprint_is_stable_and_duplicate_import_blocks_new_name(tmp_path):
    payload = exported_plan_payload(tmp_path, name="Portable Plan")
    changed_timestamp = deepcopy(payload)
    changed_timestamp["exported_at"] = "2099-01-01T00:00:00+00:00"
    reordered_plan = deepcopy(payload)
    reordered_plan["plan"] = dict(reversed(list(reordered_plan["plan"].items())))
    changed_content = deepcopy(payload)
    changed_content["versions"][0]["config_snapshot"] = changed_content["versions"][0][
        "config_snapshot"
    ].replace("2234.00", "2235.00", 1)

    assert portable_content_fingerprint(payload) == payload["portable_content_fingerprint"]
    assert portable_content_fingerprint(changed_timestamp) == portable_content_fingerprint(
        payload
    )
    assert portable_content_fingerprint(reordered_plan) == portable_content_fingerprint(payload)
    assert portable_content_fingerprint(changed_content) != portable_content_fingerprint(payload)

    path = write_payload(tmp_path, payload, "portable.json")
    target = PlanHistoryService(tmp_path / "portable-target.sqlite")
    target.import_plan_json(path)
    with pytest.raises(ValueError, match="already been imported"):
        target.import_plan_json(path, new_name="Different Name")


@pytest.mark.parametrize(
    ("section", "field", "value", "message"),
    [
        ("forecast_snapshots", "ending_savings", "999999.99", "portable content"),
        ("forecast_snapshots", "starting_debt", "1.00", "portable content"),
        ("forecast_snapshots", "total_projected_interest", "1.00", "portable content"),
        ("forecast_snapshots", "forecast_fingerprint", "bad", "portable content"),
        ("forecast_snapshots", "pay_period_count", 999, "portable content"),
        ("forecast_snapshots", "warning_count", 999, "portable content"),
        ("debt_snapshots", "ending_balance", "999.99", "portable content"),
        ("debt_snapshots", "interest_charged", "999.99", "portable content"),
        ("savings_snapshots", "withdrawal", "1.00", "portable content"),
        ("savings_snapshots", "ending_savings", "999.99", "portable content"),
        ("forecast_periods", "checking_remaining", "1.00", "portable content"),
    ],
)
def test_import_rejects_tampered_payload_without_inserting_rows(
    tmp_path,
    section,
    field,
    value,
    message,
):
    payload = exported_plan_payload(tmp_path, name=f"Tamper {section} {field}")
    payload[section][0][field] = value
    path = write_payload(tmp_path, payload, "tampered.json")
    target = PlanHistoryService(tmp_path / "tampered-target.sqlite")

    with pytest.raises(ValueError, match=message):
        target.import_plan_json(path)

    assert target.list_plans() == []


@pytest.mark.parametrize(
    ("mutate", "message"),
    [
        (
            lambda payload: payload["debt_snapshots"].__setitem__(
                0,
                {**payload["debt_snapshots"][0], "forecast_period_id": 999999},
            ),
            "orphan",
        ),
        (
            lambda payload: payload["savings_snapshots"].__setitem__(
                0,
                {**payload["savings_snapshots"][0], "forecast_period_id": 999999},
            ),
            "orphan",
        ),
        (
            lambda payload: payload["forecast_periods"].__setitem__(
                1,
                {
                    **payload["forecast_periods"][1],
                    "forecast_snapshot_id": payload["forecast_periods"][0][
                        "forecast_snapshot_id"
                    ],
                    "sequence_number": payload["forecast_periods"][0][
                        "sequence_number"
                    ],
                },
            ),
            "duplicate sequence",
        ),
        (
            lambda payload: payload["debt_snapshots"].__setitem__(
                1,
                {
                    **payload["debt_snapshots"][1],
                    "forecast_period_id": payload["debt_snapshots"][0][
                        "forecast_period_id"
                    ],
                    "debt_identifier": payload["debt_snapshots"][0][
                        "debt_identifier"
                    ],
                },
            ),
            "duplicate debt",
        ),
        (
            lambda payload: payload["actual_entries"].append(
                {
                    "id": 999,
                    "plan_id": payload["plan"]["id"],
                    "entry_date": "2026-07-17",
                    "entry_type": "not-valid",
                    "debt_identifier": None,
                    "amount": "1.00",
                    "category": "",
                    "description": "",
                    "source": "test",
                    "created_at": "2026-07-17T00:00:00+00:00",
                    "corrected_entry_id": None,
                    "note": "",
                    "forecast_period_id": None,
                    "match_method": "unmatched",
                    "matched_at": None,
                    "manual_override": 0,
                }
            ),
            "not-valid",
        ),
        (
            lambda payload: payload["actual_entries"].append(
                {
                    "id": 999,
                    "plan_id": payload["plan"]["id"],
                    "entry_date": "2026-07-17",
                    "entry_type": "income_received",
                    "debt_identifier": None,
                    "amount": "-1.00",
                    "category": "",
                    "description": "bad reversal",
                    "source": "test",
                    "created_at": "2026-07-17T00:00:00+00:00",
                    "corrected_entry_id": 123456,
                    "note": "",
                    "forecast_period_id": None,
                    "match_method": "unmatched",
                    "matched_at": None,
                    "manual_override": 0,
                }
            ),
            "invalid reversal",
        ),
    ],
)
def test_import_rejects_recomputed_tampering_and_rolls_back(tmp_path, mutate, message):
    payload = exported_plan_payload(tmp_path, name=f"Strict {message}")
    mutate(payload)
    for snapshot in payload["forecast_snapshots"]:
        snapshot["history_fingerprint"] = PlanHistoryService(
            tmp_path / "unused.sqlite"
        )._importer().snapshot_history_fingerprint(payload, int(snapshot["id"]))
    payload["portable_content_fingerprint"] = portable_content_fingerprint(payload)
    path = write_payload(tmp_path, payload, "strict-tamper.json")
    target = PlanHistoryService(tmp_path / f"strict-target-{message}.sqlite")

    with pytest.raises((ValueError, KeyError), match=message):
        target.import_plan_json(path)

    assert target.list_plans() == []


def test_import_rejects_invalid_current_version_rules(tmp_path):
    payload = exported_plan_payload(tmp_path, name="Current Rule Plan")
    no_active = deepcopy(payload)
    no_active["versions"][0]["active"] = False
    no_active["portable_content_fingerprint"] = portable_content_fingerprint(no_active)
    multiple_active = deepcopy(payload)
    second = deepcopy(multiple_active["versions"][0])
    second["id"] = 999
    second["version_number"] = 2
    second["active"] = True
    multiple_active["versions"].append(second)
    multiple_active["portable_content_fingerprint"] = portable_content_fingerprint(
        multiple_active
    )
    mismatch = deepcopy(payload)
    mismatch["plan"]["current_version_id"] = 999
    mismatch["portable_content_fingerprint"] = portable_content_fingerprint(mismatch)

    target = PlanHistoryService(tmp_path / "current-rules.sqlite")
    for index, bad_payload in enumerate((no_active, multiple_active, mismatch), start=1):
        with pytest.raises(ValueError):
            target.import_plan_json(write_payload(tmp_path, bad_payload, f"bad-{index}.json"))

    assert target.list_plans() == []


def test_import_validation_defensive_branches(tmp_path):
    base = exported_plan_payload(tmp_path, name="Validation Branch Plan")
    first_period_id = base["forecast_periods"][0]["id"]
    base["actual_entries"] = [
        {
            "id": 1,
            "plan_id": base["plan"]["id"],
            "entry_date": "2026-07-17",
            "entry_type": "income_received",
            "debt_identifier": None,
            "amount": "2234.00",
            "category": "",
            "description": "income",
            "source": "test",
            "created_at": "2026-07-17T00:00:00+00:00",
            "corrected_entry_id": None,
            "note": "",
            "forecast_period_id": first_period_id,
            "match_method": "date_window",
            "matched_at": "2026-07-17T00:00:00+00:00",
            "manual_override": 0,
        }
    ]
    base["balance_observations"] = [
        {
            "id": 1,
            "plan_id": base["plan"]["id"],
            "observation_date": "2026-07-17",
            "observation_type": "savings_balance_observation",
            "debt_identifier": None,
            "balance": "1715.00",
            "source": "test",
            "note": "",
            "created_at": "2026-07-17T00:00:00+00:00",
            "forecast_period_id": first_period_id,
            "match_method": "date_window",
        }
    ]
    base["portable_content_fingerprint"] = portable_content_fingerprint(base)
    validator = PlanHistoryService(tmp_path / "validator.sqlite")

    def refresh(payload):
        for snapshot in payload.get("forecast_snapshots", []):
            snapshot["history_fingerprint"] = validator._importer().snapshot_history_fingerprint(
                payload,
                int(snapshot["id"]),
            )
        payload["portable_content_fingerprint"] = portable_content_fingerprint(payload)
        return payload

    cases = []
    missing = deepcopy(base)
    missing.pop("debt_snapshots")
    cases.append((missing, "missing required"))
    future = deepcopy(base)
    future["source_schema_version"] = 999
    cases.append((refresh(future), "future schema"))
    no_versions = deepcopy(base)
    no_versions["versions"] = []
    cases.append((refresh(no_versions), "does not contain any versions"))
    duplicate_version_id = deepcopy(base)
    duplicate_version_id["versions"].append(
        {**duplicate_version_id["versions"][0], "version_number": 2}
    )
    cases.append((refresh(duplicate_version_id), "duplicate identifiers"))
    duplicate_version_number = deepcopy(base)
    duplicate = {**duplicate_version_number["versions"][0], "id": 999}
    duplicate_version_number["versions"].append(duplicate)
    cases.append((refresh(duplicate_version_number), "duplicate version numbers"))
    duplicate_snapshot = deepcopy(base)
    duplicate_snapshot["forecast_snapshots"].append(deepcopy(duplicate_snapshot["forecast_snapshots"][0]))
    cases.append((refresh(duplicate_snapshot), "duplicate identifiers"))
    orphan_snapshot = deepcopy(base)
    orphan_snapshot["forecast_snapshots"][0]["plan_version_id"] = 999
    cases.append((refresh(orphan_snapshot), "orphan plan_version_id"))
    duplicate_period_id = deepcopy(base)
    duplicate_period_id["forecast_periods"][1]["id"] = duplicate_period_id["forecast_periods"][0]["id"]
    cases.append((refresh(duplicate_period_id), "duplicate identifiers"))
    orphan_period = deepcopy(base)
    orphan_period["forecast_periods"][0]["forecast_snapshot_id"] = 999
    cases.append((refresh(orphan_period), "orphan forecast_snapshot_id"))
    duplicate_savings = deepcopy(base)
    duplicate_savings["savings_snapshots"].append(deepcopy(duplicate_savings["savings_snapshots"][0]))
    cases.append((refresh(duplicate_savings), "duplicate forecast periods"))
    missing_savings = deepcopy(base)
    missing_savings["savings_snapshots"].pop()
    cases.append((refresh(missing_savings), "one row per forecast period"))
    orphan_warning = deepcopy(base)
    orphan_warning["warnings"].append(
        {
            "id": 999,
            "forecast_snapshot_id": 999,
            "code": "X",
            "severity": "review",
            "message": "x",
            "relevant_date": None,
            "relevant_name": None,
            "suggested_action": None,
        }
    )
    cases.append((refresh(orphan_warning), "orphan forecast_snapshot_id"))
    bad_principal = deepcopy(base)
    bad_principal["debt_snapshots"][0]["principal_paid"] = "999.99"
    cases.append((refresh(bad_principal), "principal plus interest"))
    bad_tolerance = deepcopy(base)
    bad_tolerance["debt_snapshots"][0]["final_payoff_tolerance_applied"] = 1
    cases.append((refresh(bad_tolerance), "tolerance flag"))
    bad_deposit = deepcopy(base)
    bad_deposit["savings_snapshots"][0]["total_deposit"] = "999.99"
    cases.append((refresh(bad_deposit), "total_deposit"))
    bad_goal = deepcopy(base)
    bad_goal["savings_snapshots"][0]["goal_target"] = "1.00"
    cases.append((refresh(bad_goal), "active goal"))
    bad_start_date = deepcopy(base)
    bad_start_date["forecast_snapshots"][0]["forecast_start_date"] = "2026-01-01"
    cases.append((refresh(bad_start_date), "forecast_start_date"))
    bad_history_fingerprint = deepcopy(base)
    bad_history_fingerprint["forecast_snapshots"][0]["history_fingerprint"] = "bad"
    bad_history_fingerprint["portable_content_fingerprint"] = portable_content_fingerprint(
        bad_history_fingerprint
    )
    cases.append((bad_history_fingerprint, "history_fingerprint"))
    duplicate_actual = deepcopy(base)
    duplicate_actual["actual_entries"].append(deepcopy(duplicate_actual["actual_entries"][0]))
    cases.append((refresh(duplicate_actual), "duplicate identifiers"))
    orphan_actual = deepcopy(base)
    orphan_actual["actual_entries"][0]["forecast_period_id"] = 999
    cases.append((refresh(orphan_actual), "orphan forecast_period_id"))
    reversal_type = deepcopy(base)
    reversal_type["actual_entries"].append(
        {
            **reversal_type["actual_entries"][0],
            "id": 999,
            "entry_type": "bill_paid",
            "corrected_entry_id": reversal_type["actual_entries"][0]["id"],
        }
    )
    cases.append((refresh(reversal_type), "reversal type"))
    orphan_observation = deepcopy(base)
    orphan_observation["balance_observations"][0]["forecast_period_id"] = 999
    cases.append((refresh(orphan_observation), "orphan forecast_period_id"))

    for payload, message in cases:
        with pytest.raises(ValueError, match=message):
            validator._importer().validate_import_payload(payload)


def test_plan_version_immutability_blocks_direct_sql_but_allows_service_workflows(tmp_path):
    service = PlanHistoryService(tmp_path / "immutable.sqlite")
    config = Config().load("config.json")
    plan = service.create_plan("Immutable Plan", config)
    original = service.list_plan_versions(plan.id)[0]

    with closing(sqlite3.connect(tmp_path / "immutable.sqlite")) as conn:
        with pytest.raises(sqlite3.IntegrityError, match="immutable"):
            conn.execute(
                "UPDATE plan_versions SET change_note = 'mutated' WHERE id = ?",
                (original.id,),
            )
        with pytest.raises(sqlite3.IntegrityError, match="immutable"):
            conn.execute("DELETE FROM plan_versions WHERE id = ?", (original.id,))

    changed = deepcopy(config)
    changed.settings.paycheck = Decimal("2300.00")
    new_version = service.save_plan_version(plan.id, changed)
    restored = service.restore_plan_version(original.id)

    assert new_version.version_number == 2
    assert restored.version_number == 3
    assert service.get_plan_version(original.id).change_note != "mutated"

    service.archive_plan(plan.id)
    service.delete_plan_permanently(plan.id, confirmation_name="Immutable Plan")
    assert service.list_plans(include_archived=True) == []


def test_explanations_and_warnings_are_structured_plain_language():
    config = Config().load("config.json")
    config.savings_plan.deadline_priority_enabled = True
    forecast = ForecastEngine(config, max_years=1).forecast()
    first_period = forecast.periods[0]

    explanation = explain_forecast_period(first_period)
    warnings = build_warnings(
        forecast,
        high_apr_debts=[("Card", Decimal("35.00"))],
    )

    assert AllocationReasonCode.DEADLINE_SAVINGS_REDIRECTION in explanation.reason_codes
    assert "$215.00" in explanation.explanation
    assert all("irresponsible" not in warning.message.casefold() for warning in warnings)
    assert any(warning.severity == DataWarningSeverity.REVIEW for warning in warnings)


def test_live_and_enabled_deadline_regressions_remain_unchanged():
    config = Config().load("config.json")
    summaries, _ = short_plan(config)
    assert summaries[0].savings_contribution == Decimal("215.00")
    assert summaries[0].snowball_payment == Decimal("215.00")
    assert summaries[0].personal_expense_reduction == Decimal("0.00")
    assert summaries[1].savings_contribution == Decimal("568.50")
    assert summaries[1].snowball_payment == Decimal("568.50")
    assert summaries[1].personal_expense_reduction == Decimal("0.00")

    enabled = Config().load("config.json")
    enabled.savings_plan.deadline_priority_enabled = True
    enabled_summaries, _ = short_plan(enabled)
    assert enabled_summaries[0].savings_contribution == Decimal("430.00")
    assert enabled_summaries[0].snowball_reduction == Decimal("215.00")
    assert enabled_summaries[0].personal_expense_reduction == Decimal("0.00")
    assert enabled_summaries[1].savings_contribution == Decimal("470.00")
    assert enabled_summaries[1].snowball_reduction == Decimal("0.00")
    assert enabled_summaries[2].planned_withdrawal_amount == Decimal("2400.00")
    assert enabled_summaries[2].snowball_payment == Decimal("215.00")
