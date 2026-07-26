from datetime import date
from decimal import Decimal

from openpyxl import load_workbook

from app.excel_writer import ExcelWriter
from app.models import (
    DebtFreeTargetResult,
    DebtFreeTargetStatus,
)


def make_target_result(
    required_extra=Decimal("150.00"),
    projected_debt_free_date=date(2026, 1, 15),
    target_met=True,
    status=DebtFreeTargetStatus.TARGET_MET,
    message="Target can be reached.",
):
    return DebtFreeTargetResult(
        target_date=date(2026, 1, 15),
        required_extra_per_paycheck=required_extra,
        projected_debt_free_date=projected_debt_free_date,
        target_met=target_met,
        total_interest=Decimal("25.00"),
        total_snowball_paid=Decimal("1000.00"),
        ending_debt=Decimal("0.00") if target_met else Decimal("100.00"),
        iterations_used=18,
        lower_bound_tested=Decimal("149.99"),
        upper_bound_tested=Decimal("150.00"),
        maximum_extra_tested=Decimal("1000.00"),
        precision=Decimal("0.01"),
        calculation_status=status,
        message=message,
        baseline_debt_free_date=date(2026, 2, 12),
        baseline_total_interest=Decimal("40.00"),
        baseline_total_snowball_paid=Decimal("1015.00"),
        baseline_ending_debt=Decimal("0.00"),
    )


def test_debt_free_target_worksheet_and_dashboard_metrics(tmp_path):
    workbook_path = tmp_path / "target.xlsx"

    ExcelWriter(workbook_path).write([], None, None, make_target_result())

    workbook = load_workbook(workbook_path)
    sheet = workbook["Debt-Free Target"]
    dashboard = workbook["Dashboard"]

    assert sheet["A1"].value == "Debt-Free Target"
    assert sheet["A4"].value == "Target Date"
    assert sheet["B5"].value == 150
    assert sheet["B5"].number_format == "$#,##0.00"
    assert sheet["B4"].number_format == "mmm d, yyyy"
    assert sheet["A9"].value == "Days Accelerated"
    assert sheet["B9"].value == 28
    assert sheet["A17"].value == "Status Message"
    assert sheet.freeze_panes == "A4"
    assert sheet.auto_filter.ref == "A3:B17"
    assert dashboard["A8"].value == "Target Date"
    assert dashboard["A9"].value == "Required Extra Per Paycheck"
    assert dashboard["B9"].value == 150
    assert dashboard["A11"].value == "Target Status"
    assert dashboard["B11"].value == "Target met"
    workbook.close()


def test_unreachable_target_renders_without_misleading_required_payment(tmp_path):
    workbook_path = tmp_path / "unreachable.xlsx"
    result = make_target_result(
        required_extra=None,
        projected_debt_free_date=None,
        target_met=False,
        status=DebtFreeTargetStatus.UNREACHABLE,
        message="Target cannot be reached within the configured maximum extra payment.",
    )

    ExcelWriter(workbook_path).write([], None, None, result)

    workbook = load_workbook(workbook_path)
    assert workbook["Debt-Free Target"]["B5"].value is None
    assert workbook["Debt-Free Target"]["B17"].value.startswith("Target cannot")
    assert workbook["Dashboard"]["B9"].value is None
    assert workbook["Dashboard"]["B11"].value == "Target not reachable"
    workbook.close()


def test_disabled_target_preserves_current_workbook_behavior(tmp_path):
    workbook_path = tmp_path / "disabled.xlsx"

    ExcelWriter(workbook_path).write([], None, None, None)

    workbook = load_workbook(workbook_path)
    assert "Debt-Free Target" not in workbook.sheetnames
    assert workbook["Dashboard"]["A1"].value == "DebtPilot Dashboard"
    workbook.close()
